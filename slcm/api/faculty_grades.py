# Copyright (c) 2026, Nishanth and contributors
# For license information, please see license.txt

"""Faculty portal Grades page (www/faculty-portal/marks).

A faculty version of the desk Course Results page, limited to:
  * course offerings the faculty is assigned to (Course Offering.faculty),
  * the students enrolled in that offering,
  * what Access Result Settings allows for the exam plan + course
    (view / edit access and deadlines, the course lock, masking, grade
    override, auto-generate, grade report, visible exam types),
  * the student's marks status (Submitted / Locked rows are read-only).

Marks are stored and graded by the desk module's own functions
(write_marks_entry, write_improvement_marks, _recalculate_student_marks), so
both pages calculate identically.
"""

import base64
import io
import json

import frappe
from frappe import _
from frappe.utils import cint, flt

from slcm.slcm.page.examination_result.examination_result import (
	_get_marks_columns,
	_get_reexam_columns,
	_recalculate_student_marks,
	store_marks_entry,
	write_improvement_marks,
	write_marks_entry,
)
from slcm.slcm.utils.result_access import (
	PAID_STATUSES,
	STATUS_DRAFT,
	STATUS_LOCKED,
	STATUS_SUBMITTED,
	arrear_badge,
	assert_course_editable,
	assert_student_editable,
	edit_block_reason,
	get_access,
	get_max_marks,
	get_reexam_attempts,
	get_valid_grades,
	is_improved,
	parse_marks,
	validate_grade,
	view_block_reason,
)
from slcm.utils.faculty_portal import get_faculty_name

OFFERING_STATUSES = ("Active", "Open")


# ---------------------------------------------------------------------------
# Context: offering + exam plan + access, all checked
# ---------------------------------------------------------------------------

def _require_faculty():
	if frappe.session.user == "Guest":
		frappe.throw(_("Please log in."), frappe.PermissionError)
	faculty = get_faculty_name()
	if not faculty:
		frappe.throw(_("No faculty record is linked to your account."), frappe.PermissionError)
	return faculty


def get_faculty_offerings(faculty):
	"""Course offerings this faculty is assigned to, for the course selector."""
	offerings = frappe.get_all(
		"Course Offering",
		filters={"faculty": faculty, "status": ["in", OFFERING_STATUSES]},
		fields=["name", "course_title", "course_name", "program", "term_name", "academic_year"],
		order_by="academic_year desc, course_name asc",
		ignore_permissions=True,
	)
	programmes = {
		p.name: p.program_name or p.name
		for p in frappe.get_all(
			"Programme",
			filters={"name": ["in", list({o.program for o in offerings if o.program})]},
			fields=["name", "program_name"],
			ignore_permissions=True,
		)
	} if offerings else {}
	for o in offerings:
		o.programme_name = programmes.get(o.program) or o.program or ""
	return offerings


def _exam_plans_for(course, offering):
	"""Exam plans with a Course Schema Assignment for this course — the ones
	mapped to this offering first, then Active before Inactive, newest first."""
	return frappe.db.sql(
		"""
		SELECT csa.exam_plan, csa.evaluation_schema, csa.grade_schema, csa.course_offering,
		       ep.exam_name, ep.status AS plan_status, ep.term
		FROM `tabCourse Schema Assignment` csa
		JOIN `tabExam Plan` ep ON ep.name = csa.exam_plan
		WHERE csa.course = %(course)s
		ORDER BY (csa.course_offering = %(offering)s) DESC,
		         FIELD(ep.status, 'Active', 'Inactive') ASC, ep.creation DESC
		""",
		{"course": course, "offering": offering},
		as_dict=True,
	)


def _context(course_offering, exam_plan=None):
	"""Everything an endpoint needs, after checking the offering belongs to
	the logged-in faculty. Throws PermissionError otherwise."""
	faculty = _require_faculty()
	co = frappe.db.get_value(
		"Course Offering",
		course_offering,
		["name", "course_title", "course_name", "faculty", "program", "term_name", "academic_year"],
		as_dict=True,
	)
	# Faculty names can be autonumbered (int 1) while the link stores "1"
	if not co or str(co.faculty or "") != str(faculty):
		frappe.throw(_("You are not assigned to this course offering."), frappe.PermissionError)

	plans = _exam_plans_for(co.course_title, co.name)
	plan = None
	if exam_plan:
		plan = next((p for p in plans if p.exam_plan == exam_plan), None)
		if not plan:
			frappe.throw(_("Exam plan {0} is not set up for this course.").format(exam_plan))
	elif plans:
		plan = plans[0]

	ctx = frappe._dict(
		faculty=faculty,
		offering=co,
		course=co.course_title,
		plans=plans,
		plan=plan,
		exam_plan=plan.exam_plan if plan else None,
		evaluation_schema=plan.evaluation_schema if plan else None,
		grade_schema=plan.grade_schema if plan else None,
	)
	ctx.access = get_access(ctx.exam_plan, ctx.course)
	return ctx


def _require_plan(ctx):
	if not ctx.exam_plan:
		frappe.throw(_("No exam plan is set up for this course yet. Please contact the examination office."))


def _require_view(ctx):
	_require_plan(ctx)
	reason = view_block_reason(ctx.access)
	if reason:
		frappe.throw(reason, frappe.PermissionError)


def _require_edit(ctx):
	_require_view(ctx)
	assert_course_editable(ctx.exam_plan, ctx.course, for_faculty=True, access=ctx.access)


def _enrolled_students(offering):
	"""Students enrolled in this course offering."""
	return frappe.db.sql_list(
		"""
		SELECT DISTINCT se.student
		FROM `tabStudent Enrollment` se
		JOIN `tabStudent Enrollment Course` sec ON sec.parent = se.name
		WHERE sec.course_offering = %(offering)s
		  AND sec.status = 'Enrolled' AND se.status = 'Enrolled'
		  AND se.docstatus < 2 AND IFNULL(se.student, '') != ''
		""",
		{"offering": offering},
	)


def _assert_enrolled(ctx, student):
	if student not in set(_enrolled_students(ctx.offering.name)):
		frappe.throw(_("This student is not enrolled in your course offering."), frappe.PermissionError)


def _scm(ctx, student, create=False):
	"""The student's Student Course Marks for this course + exam plan;
	created on first save (faculty never need a manual sync)."""
	name = frappe.db.get_value(
		"Student Course Marks",
		{"course": ctx.course, "exam_plan": ctx.exam_plan, "student": student},
		"name",
	)
	if name or not create:
		return name
	doc = frappe.new_doc("Student Course Marks")
	doc.course = ctx.course
	doc.course_offering = ctx.offering.name
	doc.exam_plan = ctx.exam_plan
	doc.student = student
	doc.evaluation_schema = ctx.evaluation_schema or ""
	doc.status = STATUS_DRAFT
	doc.consider_for_sgpa = 1
	doc.insert(ignore_permissions=True)
	return doc.name


def _visible_columns(ctx):
	"""Regular assessment columns, limited to the exam types Access Result
	Settings makes visible (all, when none are set)."""
	if not ctx.evaluation_schema:
		return []
	cols = _get_marks_columns(ctx.evaluation_schema)
	visible = set(ctx.access.visible_exams or [])
	if visible:
		cols = [c for c in cols if c.assessment_type in visible]
	return cols


def _column_key(col):
	return f"{col.get('component') or ''}|{col.get('assessment_type') or ''}"


def _is_project(col):
	return (col.get("type_name") or col.get("label") or "").lower() == "project"


def _student_rows(students, masked):
	"""Display info per student, masked when Access Result Settings says so."""
	info = {
		s.name: s
		for s in frappe.get_all(
			"Student Master",
			filters={"name": ["in", students]},
			fields=["name", "first_name", "last_name", "registration_id", "passport_size_photo"],
			ignore_permissions=True,
		)
	} if students else {}
	ordered = sorted(students, key=lambda s: (info.get(s) or frappe._dict()).get("registration_id") or s)
	rows = []
	for idx, sid in enumerate(ordered, start=1):
		s = info.get(sid) or frappe._dict()
		if masked:
			rows.append({"student": sid, "student_name": f"Student {idx}", "reg_id": "—", "photo": ""})
		else:
			rows.append({
				"student": sid,
				"student_name": " ".join(filter(None, [s.get("first_name"), s.get("last_name")])) or sid,
				"reg_id": s.get("registration_id") or sid,
				"photo": s.get("passport_size_photo") or "",
			})
	return rows


def _marks_by_student(ctx, students):
	"""{student: {scm, status, total, grade, ..., entries: {comp|atype: {marks, revaluation_marks}}}}"""
	if not students or not ctx.exam_plan:
		return {}
	rows = frappe.db.sql(
		"""
		SELECT name, student, status, total_marks, grade, updated_final_marks, updated_grade,
		       re_exam_grade, improvement_marks, improvement_grade, improvement_applied,
		       attendance_status, remark
		FROM `tabStudent Course Marks`
		WHERE course = %(course)s AND exam_plan = %(plan)s AND student IN %(students)s
		""",
		{"course": ctx.course, "plan": ctx.exam_plan, "students": tuple(students)},
		as_dict=True,
	)
	out = {r.student: r for r in rows}
	attempts = get_reexam_attempts(ctx.course, students)
	for r in rows:
		r.entries = {}
		r.arrear_marker = arrear_badge(attempts.get(r.student, 0))
		r.improved = is_improved(r)
		r.status = r.status or STATUS_DRAFT
	if rows:
		for e in frappe.db.sql(
			"""
			SELECT parent, component, assessment_type, marks, revaluation_marks
			FROM `tabStudent Marks Entry`
			WHERE parent IN %(names)s
			""",
			{"names": tuple(r.name for r in rows)},
			as_dict=True,
		):
			owner = next((r for r in rows if r.name == e.parent), None)
			if owner is not None:
				owner.entries[f"{e.component or ''}|{e.assessment_type or ''}"] = {
					"marks": e.marks, "revaluation_marks": e.revaluation_marks,
				}
	return out


def _failed_grades(grade_schema):
	if not grade_schema:
		return set()
	return set(frappe.db.sql_list(
		"SELECT grade FROM `tabGrading Schema Component` WHERE parent = %s AND failed = 1",
		grade_schema,
	))


def _stats(marks, students, failed):
	graded = passed = failed_n = 0
	dist = {}
	totals = []
	for sid in students:
		m = marks.get(sid)
		grade = (m.grade or "").strip() if m and m.entries else ""
		if grade:
			graded += 1
			dist[grade] = dist.get(grade, 0) + 1
			if grade in failed:
				failed_n += 1
			else:
				passed += 1
		if m and m.total_marks is not None and m.entries:
			totals.append(flt(m.total_marks))
	submitted = sum(1 for sid in students if marks.get(sid) and marks[sid].status == STATUS_SUBMITTED)
	return {
		"total": len(students),
		"graded": graded,
		"not_graded": len(students) - graded,
		"passed": passed,
		"failed": failed_n,
		"submitted": submitted,
		"avg_marks": round(sum(totals) / len(totals), 2) if totals else None,
		"grade_dist": [{"grade": g, "count": c} for g, c in sorted(dist.items(), key=lambda x: -x[1])],
	}


# ---------------------------------------------------------------------------
# Grade sheet (main tab)
# ---------------------------------------------------------------------------

@frappe.whitelist()
def get_grade_sheet(course_offering, exam_plan=None):
	ctx = _context(course_offering, exam_plan)
	access = ctx.access
	info = {
		"course_offering": ctx.offering.name,
		"course": ctx.course,
		"course_name": ctx.offering.course_name or ctx.course,
		"term": ctx.offering.term_name or "",
		"academic_year": ctx.offering.academic_year or "",
		"exam_plans": [
			{"name": p.exam_plan, "label": p.exam_name or p.exam_plan, "status": p.plan_status}
			for p in ctx.plans
		],
		"exam_plan": ctx.exam_plan,
		"evaluation_schema": ctx.evaluation_schema or "",
		"grade_schema": ctx.grade_schema or "",
		"view_block": view_block_reason(access) if ctx.exam_plan else None,
		"edit_block": edit_block_reason(access, for_faculty=True) if ctx.exam_plan else None,
		"edit_deadline": frappe.utils.format_datetime(access.edit_deadline, "dd MMM yyyy, hh:mm a") if access.edit_deadline else "",
		"masked": bool(access.mask_student_info),
		"can_override_grade": bool(access.edit_grade_access),
		"can_auto_generate": bool(access.auto_generate_grade_access),
		"can_grade_report": bool(access.generate_grade_report),
		"visible_exams": access.visible_exams,
		"grades": sorted(get_valid_grades(ctx.grade_schema)),
		"failed_grades": sorted(_failed_grades(ctx.grade_schema)),
	}
	if not ctx.exam_plan or info["view_block"]:
		return {"info": info, "columns": [], "students": [], "stats": None}

	columns = [
		{
			"key": _column_key(c),
			"component": c.component,
			"assessment_type": c.assessment_type,
			"label": c.label or c.component_name or c.component,
			"type_name": c.type_name or "",
			"max_marks": flt(c.maximum_marks),
			"is_project": _is_project(c),
		}
		for c in _visible_columns(ctx)
	]

	students = _enrolled_students(ctx.offering.name)
	marks = _marks_by_student(ctx, students)
	rows = []
	for r in _student_rows(students, info["masked"]):
		m = marks.get(r["student"])
		r.update({
			"status": m.status if m else STATUS_DRAFT,
			# No marks entered yet → not graded (the recalculation would say 0 / F)
			"total": m.total_marks if m and m.entries else None,
			"grade": (m.grade or "") if m and m.entries else "",
			"updated_final_marks": m.updated_final_marks if m else None,
			"updated_grade": (m.updated_grade or "") if m else "",
			"attendance_status": (m.attendance_status or "") if m else "",
			"remark": (m.remark or "") if m else "",
			"arrear_marker": m.arrear_marker if m else "",
			"improved": bool(m.improved) if m else False,
			"entries": m.entries if m else {},
		})
		rows.append(r)

	stats = _stats(marks, students, set(info["failed_grades"])) if info["can_grade_report"] else None
	return {"info": info, "columns": columns, "students": rows, "stats": stats}


@frappe.whitelist()
def save_mark(course_offering, exam_plan, student, component, assessment_type, value, field="marks"):
	"""Inline marks entry for one regular component."""
	ctx = _context(course_offering, exam_plan)
	_require_edit(ctx)
	_assert_enrolled(ctx, student)
	if field not in ("marks", "revaluation_marks"):
		frappe.throw(_("Invalid field."))
	key = f"{component or ''}|{assessment_type or ''}"
	if key not in {_column_key(c) for c in _visible_columns(ctx)}:
		frappe.throw(_("This assessment is not open for entry."), frappe.PermissionError)
	scm = _scm(ctx, student, create=True)
	assert_student_editable(scm, for_faculty=True)
	result = write_marks_entry(scm, ctx.course, ctx.exam_plan, component, assessment_type, field, value)
	return _with_badges(ctx, student, result)


@frappe.whitelist()
def save_grade(course_offering, exam_plan, student, grade):
	"""Grade override — only when Access Result Settings has Edit Grade Access."""
	ctx = _context(course_offering, exam_plan)
	_require_edit(ctx)
	if not ctx.access.edit_grade_access:
		frappe.throw(_("Grade override is not enabled for this course."), frappe.PermissionError)
	_assert_enrolled(ctx, student)
	scm = _scm(ctx, student, create=True)
	assert_student_editable(scm, for_faculty=True)
	grade = validate_grade(grade, ctx.grade_schema)
	frappe.db.set_value("Student Course Marks", scm, "grade", grade)
	frappe.db.commit()
	return {"grade": grade}


@frappe.whitelist()
def save_remark(course_offering, exam_plan, student, remark):
	ctx = _context(course_offering, exam_plan)
	_require_edit(ctx)
	_assert_enrolled(ctx, student)
	scm = _scm(ctx, student, create=True)
	assert_student_editable(scm, for_faculty=True)
	frappe.db.set_value("Student Course Marks", scm, "remark", (remark or "").strip())
	frappe.db.commit()
	return {"ok": True}


@frappe.whitelist()
def auto_generate_grades(course_offering, exam_plan):
	"""Recalculate total + grade for every Draft student of the offering —
	only when Access Result Settings has Auto Generate Grade Access."""
	ctx = _context(course_offering, exam_plan)
	_require_edit(ctx)
	if not ctx.access.auto_generate_grade_access:
		frappe.throw(_("Auto grade generation is not enabled for this course."), frappe.PermissionError)
	students = _enrolled_students(ctx.offering.name)
	done = 0
	for row in frappe.get_all(
		"Student Course Marks",
		filters={"course": ctx.course, "exam_plan": ctx.exam_plan, "student": ["in", students or ["__none__"]]},
		fields=["name", "status"],
		ignore_permissions=True,
	):
		if (row.status or STATUS_DRAFT) != STATUS_DRAFT:
			continue
		_recalculate_student_marks(row.name, ctx.course, ctx.exam_plan)
		done += 1
	return {"updated": done}


@frappe.whitelist()
def set_status(course_offering, exam_plan, status, students):
	"""Mark as Submitted / Mark as Draft for the faculty's own students.
	Locked rows belong to the examination office and are skipped."""
	ctx = _context(course_offering, exam_plan)
	_require_edit(ctx)
	if status not in (STATUS_SUBMITTED, STATUS_DRAFT):
		frappe.throw(_("Faculty can only set Submitted or Draft."))
	if isinstance(students, str):
		students = json.loads(students or "[]")
	enrolled = set(_enrolled_students(ctx.offering.name))
	students = [s for s in (students or []) if s in enrolled]
	if not students:
		frappe.throw(_("Select at least one student."))

	updated = skipped = 0
	for sid in students:
		scm = _scm(ctx, sid, create=True)
		current = frappe.db.get_value("Student Course Marks", scm, "status") or STATUS_DRAFT
		if current == STATUS_LOCKED:
			skipped += 1
			continue
		if current != status:
			frappe.db.set_value("Student Course Marks", scm, "status", status)
			updated += 1
	frappe.db.commit()
	return {"updated": updated, "skipped_locked": skipped, "status": status}


def _with_badges(ctx, student, result):
	result = dict(result or {})
	result["arrear_marker"] = arrear_badge(get_reexam_attempts(ctx.course, [student]).get(student, 0))
	return result


# ---------------------------------------------------------------------------
# Excel template / import
# ---------------------------------------------------------------------------

def _template_columns(ctx):
	"""[(header, key, field, max)] — Project components get Marks + Deduction."""
	out = []
	for c in _visible_columns(ctx):
		label = c.label or c.component_name or c.component
		key = _column_key(c)
		mx = flt(c.maximum_marks)
		if _is_project(c):
			out.append((f"{label} - Marks (Max: {mx:g})", key, "marks", mx))
			out.append((f"{label} - Deduction (Max: {mx:g})", key, "revaluation_marks", mx))
		else:
			out.append((f"{label} (Max: {mx:g})", key, "marks", mx))
	return out


@frappe.whitelist()
def download_template(course_offering, exam_plan):
	"""Excel template for the offering: one row per enrolled student with
	their current marks. Column A (hidden) is the row key used on import."""
	import openpyxl
	from openpyxl.styles import Alignment, Font, PatternFill

	ctx = _context(course_offering, exam_plan)
	_require_view(ctx)
	cols = _template_columns(ctx)
	with_grade = bool(ctx.access.edit_grade_access)

	students = _enrolled_students(ctx.offering.name)
	marks = _marks_by_student(ctx, students)
	rows = _student_rows(students, bool(ctx.access.mask_student_info))

	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "Marks"
	# The grade column starts blank: only a grade typed in is an override. A
	# pre-filled current grade would re-apply a stale grade over the one
	# recalculated from the uploaded marks.
	headers = ["Key", "Registration ID", "Student Name"] + [c[0] for c in cols] + (["Grade Override (optional)"] if with_grade else [])
	# Row 2 carries machine keys for each column (hidden) so renamed headers
	# can't silently map marks to the wrong component.
	keys = ["key", "reg", "name"] + [f"{c[1]}#{c[2]}" for c in cols] + (["grade"] if with_grade else [])
	for i, h in enumerate(headers, start=1):
		cell = ws.cell(row=1, column=i, value=h)
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = PatternFill("solid", fgColor="920C24")
		cell.alignment = Alignment(horizontal="center", wrap_text=True)
		ws.cell(row=2, column=i, value=keys[i - 1])
	ws.row_dimensions[2].hidden = True
	ws.column_dimensions["A"].hidden = True
	ws.column_dimensions["B"].width = 18
	ws.column_dimensions["C"].width = 30
	for i in range(4, len(headers) + 1):
		ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18
	ws.freeze_panes = "D3"

	for r_idx, r in enumerate(rows, start=3):
		m = marks.get(r["student"])
		ws.cell(row=r_idx, column=1, value=r["student"])
		ws.cell(row=r_idx, column=2, value=r["reg_id"])
		ws.cell(row=r_idx, column=3, value=r["student_name"])
		for c_idx, (_h, key, field, _mx) in enumerate(cols, start=4):
			val = ((m.entries.get(key) or {}).get(field)) if m else None
			ws.cell(row=r_idx, column=c_idx, value=val)

	buf = io.BytesIO()
	wb.save(buf)
	frappe.response.filename = f"marks_{ctx.offering.name}_{ctx.exam_plan}.xlsx".replace(" ", "_")
	frappe.response.filecontent = buf.getvalue()
	frappe.response.type = "binary"


@frappe.whitelist()
def import_marks(course_offering, exam_plan, filedata, filename="", apply=0):
	"""Validate an uploaded template (base64) and, with apply=1, save it.

	Checks: file type and template format, every student is enrolled in the
	offering and not Submitted/Locked, marks are numbers within 0..max,
	grades are in the grading schema. Nothing is saved unless every row is
	valid. Returns {ok, errors: [{row, student, column, message}], summary}."""
	ctx = _context(course_offering, exam_plan)
	_require_edit(ctx)
	if not (filename or "").lower().endswith(".xlsx"):
		frappe.throw(_("Please upload the .xlsx template downloaded from this page."))

	import openpyxl

	try:
		raw = base64.b64decode((filedata or "").split(",")[-1])
		ws = openpyxl.load_workbook(io.BytesIO(raw), data_only=True).active
	except Exception:
		frappe.throw(_("The file could not be read as an Excel (.xlsx) workbook."))

	keys = [str(ws.cell(row=2, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
	headers = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
	if not keys or keys[0] != "key":
		frappe.throw(_("This is not the marks template from this page. Download the template and fill that in."))

	allowed = {f"{k}#{f}": mx for (_h, k, f, mx) in _template_columns(ctx)}
	grade_col = keys.index("grade") if "grade" in keys else None
	if grade_col is not None and not ctx.access.edit_grade_access:
		grade_col = None  # template from before access was removed — ignore grades
	mark_cols = [(i, k) for i, k in enumerate(keys) if "#" in k]
	unknown = [headers[i] or k for i, k in mark_cols if k not in allowed]
	if unknown:
		frappe.throw(_("These columns are not open for entry: {0}").format(", ".join(unknown)))

	enrolled = set(_enrolled_students(ctx.offering.name))
	statuses = {
		r.student: (r.status or STATUS_DRAFT)
		for r in frappe.get_all(
			"Student Course Marks",
			filters={"course": ctx.course, "exam_plan": ctx.exam_plan},
			fields=["student", "status"],
			ignore_permissions=True,
		)
	}
	reg_to_student = {
		r.registration_id: r.name
		for r in frappe.get_all(
			"Student Master",
			filters={"name": ["in", list(enrolled) or ["__none__"]]},
			fields=["name", "registration_id"],
			ignore_permissions=True,
		) if r.registration_id
	}

	errors, pending, seen = [], {}, set()
	for row_no in range(3, ws.max_row + 1):
		row = [ws.cell(row=row_no, column=c + 1).value for c in range(len(keys))]
		if all(v is None or str(v).strip() == "" for v in row):
			continue
		sid = str(row[0] or "").strip()
		reg = str(row[1] or "").strip()
		who = reg if reg and reg != "—" else (row[2] or sid)
		if sid not in enrolled:
			sid = reg_to_student.get(reg, sid)
		if sid not in enrolled:
			errors.append({"row": row_no, "student": who, "column": "", "message": _("Student is not enrolled in this course offering.")})
			continue
		if sid in seen:
			errors.append({"row": row_no, "student": who, "column": "", "message": _("Student appears more than once in the file.")})
			continue
		seen.add(sid)
		if statuses.get(sid, STATUS_DRAFT) != STATUS_DRAFT:
			errors.append({"row": row_no, "student": who, "column": "", "message": _("Marks are {0} and can't be changed.").format(statuses[sid])})
			continue

		values = []
		for i, k in mark_cols:
			cell = row[i]
			if cell is None or str(cell).strip() == "":
				continue
			try:
				values.append((k, parse_marks(cell, allowed[k], label=headers[i] or k)))
			except frappe.ValidationError as e:
				errors.append({"row": row_no, "student": who, "column": headers[i], "message": frappe.utils.strip_html(str(e))})
				frappe.clear_messages()
		grade = None
		if grade_col is not None and row[grade_col] not in (None, ""):
			try:
				grade = validate_grade(str(row[grade_col]), ctx.grade_schema)
			except frappe.ValidationError as e:
				errors.append({"row": row_no, "student": who, "column": "Grade", "message": frappe.utils.strip_html(str(e))})
				frappe.clear_messages()
		if values or grade:
			pending[sid] = (values, grade)

	summary = {"students": len(pending), "values": sum(len(v[0]) for v in pending.values())}
	if errors or not cint(apply):
		return {"ok": not errors, "errors": errors, "summary": summary, "applied": False}

	for sid, (values, grade) in pending.items():
		scm = _scm(ctx, sid, create=True)
		for k, val in values:
			comp_atype, field = k.split("#")
			component, assessment_type = comp_atype.split("|")
			store_marks_entry(scm, component, assessment_type, field, val)
		_recalculate_student_marks(scm, ctx.course, ctx.exam_plan)
		if grade:
			frappe.db.set_value("Student Course Marks", scm, "grade", grade)
	frappe.db.commit()
	return {"ok": True, "errors": [], "summary": summary, "applied": True}


# ---------------------------------------------------------------------------
# Re-exam / improvement exam tab
# ---------------------------------------------------------------------------

def _paid_registrations(ctx, doctype):
	"""Paid registrations of this course + exam plan for students enrolled in
	this offering. Enrolment decides, not the registration's own
	course_offering field — that isn't reliable (it can point at another
	section, or even another course)."""
	enrolled = set(_enrolled_students(ctx.offering.name))
	rows = frappe.get_all(
		doctype,
		filters={
			"course": ctx.course,
			"exam_plan": ctx.exam_plan,
			"status": "Registered",
			"payment_status": ["in", PAID_STATUSES],
		},
		fields=["name", "student"],
		ignore_permissions=True,
	)
	return {r.student: r.name for r in rows if r.student in enrolled}


@frappe.whitelist()
def get_reexam_sheet(course_offering, exam_plan=None):
	"""Students with a paid Re Exam / Improvement Exam registration, with
	their previous result next to the new marks inputs."""
	ctx = _context(course_offering, exam_plan)
	_require_view(ctx)
	reexam_regs = _paid_registrations(ctx, "Re Exam Registration")
	improvement_regs = _paid_registrations(ctx, "Improvement Exam Registration")
	students = sorted(set(reexam_regs) | set(improvement_regs))

	reexam_cols = [
		{
			"key": _column_key(c),
			"component": c.component,
			"assessment_type": c.assessment_type,
			"label": c.label or c.component_name or c.component,
			"max_marks": flt(c.maximum_marks),
		}
		for c in (_get_reexam_columns(ctx.evaluation_schema) if ctx.evaluation_schema else [])
	]

	marks = _marks_by_student(ctx, students)
	rows = []
	for r in _student_rows(students, bool(ctx.access.mask_student_info)):
		sid = r["student"]
		m = marks.get(sid)
		r.update({
			"reexam": sid in reexam_regs,
			"improvement": sid in improvement_regs,
			"previous_total": m.total_marks if m else None,
			"previous_grade": (m.grade or "") if m else "",
			"re_exam_grade": (m.re_exam_grade or "") if m else "",
			"improvement_marks": m.improvement_marks if m and flt(m.improvement_marks) > 0 else None,
			"improvement_grade": (m.improvement_grade or "") if m else "",
			"updated_final_marks": m.updated_final_marks if m else None,
			"updated_grade": (m.updated_grade or "") if m else "",
			"arrear_marker": m.arrear_marker if m else "",
			"improved": bool(m.improved) if m else False,
			"entries": m.entries if m else {},
		})
		rows.append(r)
	return {
		"columns": reexam_cols,
		"students": rows,
		"edit_block": edit_block_reason(ctx.access, for_faculty=True),
	}


@frappe.whitelist()
def save_reexam_mark(course_offering, exam_plan, student, component, assessment_type, value):
	"""Re-exam marks — only for students with a paid re-exam registration.
	A student's Submitted status does not block this (re-exam comes after
	the regular result); the course lock and edit access still apply."""
	ctx = _context(course_offering, exam_plan)
	_require_edit(ctx)
	if student not in _paid_registrations(ctx, "Re Exam Registration"):
		frappe.throw(_("This student has no paid re-exam registration for this course."), frappe.PermissionError)
	allowed = {_column_key(c) for c in (_get_reexam_columns(ctx.evaluation_schema) if ctx.evaluation_schema else [])}
	if f"{component or ''}|{assessment_type or ''}" not in allowed:
		frappe.throw(_("This is not a re-exam component of the course."))
	scm = _scm(ctx, student)
	if not scm:
		frappe.throw(_("No previous result exists for this student."))
	if (frappe.db.get_value("Student Course Marks", scm, "status") or "") == STATUS_LOCKED:
		frappe.throw(_("This student's marks are locked."))
	result = write_marks_entry(scm, ctx.course, ctx.exam_plan, component, assessment_type, "marks", value)
	return _with_badges(ctx, student, result)


@frappe.whitelist()
def save_improvement_mark(course_offering, exam_plan, student, value):
	"""Improvement-exam marks — only for students with a paid improvement
	registration. Applied to the final result (I badge) when higher."""
	ctx = _context(course_offering, exam_plan)
	_require_edit(ctx)
	if student not in _paid_registrations(ctx, "Improvement Exam Registration"):
		frappe.throw(_("This student has no paid improvement exam registration for this course."), frappe.PermissionError)
	scm = _scm(ctx, student)
	if not scm:
		frappe.throw(_("No previous result exists for this student."))
	if (frappe.db.get_value("Student Course Marks", scm, "status") or "") == STATUS_LOCKED:
		frappe.throw(_("This student's marks are locked."))
	result = write_improvement_marks(scm, ctx.course, ctx.exam_plan, value)
	row = frappe.db.get_value(
		"Student Course Marks", scm,
		["total_marks", "updated_final_marks", "improvement_applied"], as_dict=True,
	)
	result = _with_badges(ctx, student, result)
	result["improved"] = is_improved(row)
	return result
