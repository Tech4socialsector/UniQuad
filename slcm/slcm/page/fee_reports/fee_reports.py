"""Fee Reports page — Due Wise Report and Student Wise Outstanding.

Column headers follow the office's sheets in ~/Desktop/SLCM. Students are picked with the same
filters as the Student Fee Management list (shared `_student_base`), and the Due Wise Report adds
demand-level filters: demand date, settlement date and demand type.
"""

import frappe
from frappe import _
from frappe.utils import cint, flt, getdate, today

from slcm.slcm.page.student_fee_management.student_fee_management import (
	_as_list,
	_check_access,
	_student_base,
	get_filter_options as _student_filter_options,
)

# (key, sheet header, is_amount)
DUE_WISE_COLUMNS = [
	("student_id", "Student ID", False),
	("student_name", "Student Name", False),
	("email", "Email", False),
	("programme", "Programme", False),
	("batch", "Batch", False),
	("due_type", "Due Type", False),
	("voucher_number", "Voucher Number", False),
	("fee_component", "Fee Component", False),
	("remarks", "Remarks", False),
	("original_amount", "Original Amount", True),
	("waiver_amount", "Scholarship / Waiver Amount", True),
	("penalty_amount", "Penalty Amount", True),
	("net_payable", "Net Payable", True),
	("paid_amount", "Amount Paid", True),
	("outstanding_amount", "Amount Outstanding", True),
	("due_status", "Due Status", False),
	("demand_date", "Demand Date", False),
	("due_date", "Due Date", False),
	("settlement_date", "Settlement Date", False),
]
_DUE_AMOUNT_KEYS = [key for key, _label, is_amount in DUE_WISE_COLUMNS if is_amount]

# A due is settled once nothing is outstanding and something was paid; it settled on its last payment.
_SETTLED = "IFNULL(d.outstanding_amount, 0) <= 0 AND IFNULL(d.paid_amount, 0) > 0"

DEMAND_TYPES = ("Academic", "Non Academic")


@frappe.whitelist()
def get_filter_options():
	"""The Student Fee Management filter options plus the demand types."""
	options = _student_filter_options()
	options["demand_types"] = list(DEMAND_TYPES)
	return options


def _date(value):
	return getdate(value) if value else None


def _report_filters(
	academic_year=None,
	academic_term=None,
	programme=None,
	dues_status=None,
	search=None,
	demand_type=None,
	demand_from=None,
	demand_to=None,
	settlement_from=None,
	settlement_to=None,
):
	filters = {
		"academic_year": academic_year,
		"academic_term": academic_term,
		"programme": programme,
		"dues_status": dues_status,
		"search": (search or "").strip(),
		"demand_type": [t for t in _as_list(demand_type) if t in DEMAND_TYPES],
		"demand_from": _date(demand_from),
		"demand_to": _date(demand_to),
		"settlement_from": _date(settlement_from),
		"settlement_to": _date(settlement_to),
	}
	for start, end, label in (
		("demand_from", "demand_to", _("Demand Date")),
		("settlement_from", "settlement_to", _("Settlement Date")),
	):
		if filters[start] and filters[end] and filters[start] > filters[end]:
			frappe.throw(_("{0}: the From date must be on or before the To date.").format(label))
	return filters


def _report_students_sql(filters):
	base, values = _student_base(
		filters.get("academic_year"),
		filters.get("academic_term"),
		filters.get("programme"),
		filters.get("dues_status"),
		None,
		filters.get("search"),
	)
	return f"SELECT sm.name {base}", values


def _demand_conditions(filters, values):
	"""Due-level filters of the Due Wise Report (alias d = Fee Demand)."""
	conditions = []
	if filters.get("demand_type"):
		conditions.append("d.demand_type IN %(demand_type)s")
		values["demand_type"] = tuple(filters["demand_type"])
	if filters.get("demand_from"):
		conditions.append("d.demand_date >= %(demand_from)s")
		values["demand_from"] = filters["demand_from"]
	if filters.get("demand_to"):
		conditions.append("d.demand_date <= %(demand_to)s")
		values["demand_to"] = filters["demand_to"]
	# Any settlement-date bound keeps only settled dues.
	if filters.get("settlement_from") or filters.get("settlement_to"):
		conditions.append(f"({_SETTLED})")
	if filters.get("settlement_from"):
		conditions.append("d.last_payment_date >= %(settlement_from)s")
		values["settlement_from"] = filters["settlement_from"]
	if filters.get("settlement_to"):
		conditions.append("d.last_payment_date <= %(settlement_to)s")
		values["settlement_to"] = filters["settlement_to"]
	return conditions


def _due_wise(filters, start=None, page_length=None):
	"""One row per (non-cancelled) Fee Demand of the filtered students."""
	students_sql, values = _report_students_sql(filters)
	conditions = ["d.status != 'Cancelled'", f"d.student IN ({students_sql})", *_demand_conditions(filters, values)]
	where = " AND ".join(conditions)
	totals = frappe.db.sql(
		f"""SELECT COUNT(*) AS row_count,
			{", ".join(f"SUM(IFNULL(d.{k}, 0)) AS {k}" for k in _DUE_AMOUNT_KEYS)}
		FROM `tabFee Demand` d WHERE {where}""",
		values,
		as_dict=True,
	)[0]
	limit = ""
	if page_length:
		values.update(start=cint(start), page_length=min(cint(page_length), 500))
		limit = "LIMIT %(start)s, %(page_length)s"
	rows = frappe.db.sql(
		f"""SELECT
			d.student,
			COALESCE(NULLIF(sm.registration_id, ''), sm.name) AS student_id,
			COALESCE(NULLIF(d.student_name, ''), sm.first_name) AS student_name,
			COALESCE(NULLIF(d.student_email, ''), NULLIF(sm.official_email_id, ''), sm.email) AS email,
			sm.programme_of_study AS programme, sm.batch,
			d.demand_type AS due_type, d.name AS voucher_number, d.fee_component, d.remarks,
			{", ".join(f"IFNULL(d.{k}, 0) AS {k}" for k in _DUE_AMOUNT_KEYS)},
			d.status AS due_status, d.demand_date, d.due_date,
			CASE WHEN {_SETTLED} THEN d.last_payment_date END AS settlement_date
		FROM `tabFee Demand` d
		JOIN `tabStudent Master` sm ON sm.name = d.student
		WHERE {where}
		ORDER BY sm.first_name, sm.name, d.demand_date, d.name
		{limit}""",
		values,
		as_dict=True,
	)
	return rows, totals


def _outstanding(filters, start=None, page_length=None):
	"""One row per filtered student: outstanding per fee component (every Fee Component
	is a column, as in the sheet) and the total."""
	students_sql, values = _report_students_sql(filters)
	components = frappe.get_all("Fee Component", pluck="name", order_by="name asc")

	count = frappe.db.sql(f"SELECT COUNT(*) FROM ({students_sql}) s", values)[0][0]
	limit = ""
	if page_length:
		values.update(start=cint(start), page_length=min(cint(page_length), 500))
		limit = "LIMIT %(start)s, %(page_length)s"
	students = frappe.db.sql(
		f"""SELECT sm.name AS student,
			COALESCE(NULLIF(sm.registration_id, ''), sm.name) AS student_id,
			sm.first_name AS student_name,
			COALESCE(NULLIF(sm.official_email_id, ''), sm.email) AS email,
			sm.academic_status
		FROM `tabStudent Master` sm
		WHERE sm.name IN ({students_sql})
		ORDER BY sm.first_name, sm.name
		{limit}""",
		values,
		as_dict=True,
	)

	def by_component(student_condition, params):
		return frappe.db.sql(
			f"""SELECT student, fee_component, SUM(IFNULL(outstanding_amount, 0)) AS amount
			FROM `tabFee Demand`
			WHERE status != 'Cancelled' AND student IN {student_condition}
			GROUP BY student, fee_component""",
			params,
			as_dict=True,
		)

	page_amounts = by_component("%(page_students)s", {"page_students": tuple(s.student for s in students) or ("",)})
	grid = {}
	for a in page_amounts:
		grid.setdefault(a.student, {})[a.fee_component] = flt(a.amount)
		if a.fee_component not in components:
			components.append(a.fee_component)

	for s in students:
		s["amounts"] = {c: grid.get(s.student, {}).get(c, 0) for c in components}
		s["total"] = sum(s["amounts"].values())

	# Totals across every filtered student, not just this page.
	totals = {c: 0 for c in components}
	for a in by_component(f"({students_sql})", values):
		totals[a.fee_component] = totals.get(a.fee_component, 0) + flt(a.amount)
	totals["total"] = sum(v for k, v in totals.items() if k != "total")
	return components, students, count, totals


@frappe.whitelist()
def get_due_wise_report(start=0, page_length=25, **filters):
	_check_access()
	rows, totals = _due_wise(_report_filters(**_known(filters)), start, page_length or 25)
	return {"columns": [[k, label, a] for k, label, a in DUE_WISE_COLUMNS], "rows": rows, "totals": totals}


@frappe.whitelist()
def get_outstanding_report(start=0, page_length=25, **filters):
	_check_access()
	components, rows, count, totals = _outstanding(_report_filters(**_known(filters)), start, page_length or 25)
	return {"components": components, "rows": rows, "count": count, "totals": totals}


_FILTER_KEYS = (
	"academic_year",
	"academic_term",
	"programme",
	"dues_status",
	"search",
	"demand_type",
	"demand_from",
	"demand_to",
	"settlement_from",
	"settlement_to",
)


def _known(filters):
	"""Drop request extras (cmd, csrf…) so only report filters reach _report_filters."""
	return {k: v for k, v in filters.items() if k in _FILTER_KEYS}


@frappe.whitelist()
def export_report(report, **filters):
	"""Download the whole (unpaginated) report as .xlsx in the sheet's layout."""
	_check_access()
	from io import BytesIO

	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.utils import get_column_letter

	filters = _report_filters(**_known(filters))
	wb = Workbook()
	ws = wb.active
	bold = Font(bold=True)
	header_fill = PatternFill("solid", fgColor="F2F2F2")
	amount_format = "#,##,##0.00"

	if report == "due_wise":
		ws.title = "Due Wise Report"
		rows, totals = _due_wise(filters)
		ws.append([label for _k, label, _a in DUE_WISE_COLUMNS])
		for r in rows:
			ws.append([
				(flt(r[k]) if is_amount else (r[k] if r[k] is not None else "")) for k, _label, is_amount in DUE_WISE_COLUMNS
			])
		total_row = ["Total"] + [""] * (len(DUE_WISE_COLUMNS) - 1)
		for i, (k, _label, is_amount) in enumerate(DUE_WISE_COLUMNS):
			if is_amount:
				total_row[i] = flt(totals.get(k))
		ws.append(total_row)
		amount_cols = [i + 1 for i, c in enumerate(DUE_WISE_COLUMNS) if c[2]]
		date_cols = [i + 1 for i, c in enumerate(DUE_WISE_COLUMNS) if c[0].endswith("_date")]
		filename = f"Due Wise Report {today()}.xlsx"
	elif report == "outstanding":
		ws.title = "Student wise outstanding fee"
		components, rows, _count, totals = _outstanding(filters)
		ws.append(["Sl. No.", "Student ID", "Student Name", "Student Email ID", "Academic Status"] + components + ["Total Outstanding Fee"])
		for i, r in enumerate(rows, 1):
			ws.append(
				[i, r.student_id, r.student_name or "", r.email or "", r.academic_status or ""]
				+ [flt(r["amounts"][c]) for c in components]
				+ [flt(r["total"])]
			)
		ws.append(["Total", "", "", "", ""] + [flt(totals.get(c)) for c in components] + [flt(totals["total"])])
		amount_cols = list(range(6, 6 + len(components) + 1))
		date_cols = []
		filename = f"Student wise outstanding fee report {today()}.xlsx"
	else:
		frappe.throw(_("Unknown report"))

	for cell in ws[1]:
		cell.font = bold
		cell.fill = header_fill
		cell.alignment = Alignment(wrap_text=True, vertical="center")
	for cell in ws[ws.max_row]:
		cell.font = bold
	for col in amount_cols:
		for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
			row[0].number_format = amount_format
	for col in date_cols:
		for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
			row[0].number_format = "DD-MM-YYYY"
	for i in range(1, ws.max_column + 1):
		ws.column_dimensions[get_column_letter(i)].width = 18
	ws.freeze_panes = "A2"

	out = BytesIO()
	wb.save(out)
	frappe.local.response.filename = filename
	frappe.local.response.filecontent = out.getvalue()
	frappe.local.response.type = "binary"
