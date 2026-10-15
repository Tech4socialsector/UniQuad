# Copyright (c) 2026, TFSS and contributors
# For license information, please see license.txt

"""Faculty portal Assessment Marks page (www/faculty-portal/assessment-marks).

Faculty enter 1st and 2nd submission marks for each assessment set up in the
Exam Assessment Schedule (Assessment 1, Assessment 2, Re-Exam ...). Marks are
stored per student in Student Course Marks → Assessment Marks (child table
Student Assessment Marks), one row per assessment.

Entry window, per assessment row of the schedule:
  * opens when the assessment closes — the 2nd submission deadline when set,
    else the 1st, else the assessment date;
  * closes at the row's Marks Entry Deadline (open-ended when blank).
The course LOCK (Access Result Settings) and a student's own Locked /
Submitted marks status also block entry, as on the Grades page.
"""

import base64
import io
import json
import re

import frappe
from frappe import _
from frappe.utils import cint, flt, get_datetime, now_datetime

from slcm.api.faculty_grades import (
	_assert_enrolled,
	_context,
	_enrolled_students,
	_scm,
	_student_rows,
)
from slcm.slcm.utils.result_access import (
	STATUS_DRAFT,
	STATUS_LOCKED,
	STATUS_SUBMITTED,
	get_max_marks,
	parse_marks,
)

WINDOW_OPEN = "open"
WINDOW_NOT_OPEN = "not_open"
WINDOW_CLOSED = "closed"
WINDOW_UNSCHEDULED = "unscheduled"


def _fmt(dt):
	return frappe.utils.format_datetime(dt, "d MMM yyyy, h:mm a") if dt else ""


def _schedule_items(offering):
	"""Active Exam Assessment Schedule rows for this offering. Matched on
	course + programme + term (like the student page), so a duplicate offering
	of the same course still finds its schedule."""
	return frappe.db.sql(
		"""
		SELECT item.name, item.parent AS assessment_schedule, eas.exam_plan,
		       item.assessment_label, item.component, item.assessment_type,
		       item.assessment_date, item.first_submission_deadline,
		       item.second_submission_deadline, item.marks_entry_deadline, item.remarks
		FROM `tabExam Assessment Schedule Item` item
		JOIN `tabExam Assessment Schedule` eas ON eas.name = item.parent
		WHERE item.parenttype = 'Exam Assessment Schedule'
		  AND eas.status = 'Active'
		  AND item.course = %(course)s
		  AND eas.programme = %(programme)s
		  AND eas.term = %(term)s
		ORDER BY COALESCE(item.assessment_date, item.first_submission_deadline) IS NULL,
		         COALESCE(item.assessment_date, item.first_submission_deadline), item.idx
		""",
		{"course": offering.course_title, "programme": offering.program, "term": offering.term_name},
		as_dict=True,
	)


def _window(item):
	"""(state, opens_at, closes_at) of the marks-entry window for a schedule row."""
	opens = item.second_submission_deadline or item.first_submission_deadline or item.assessment_date
	closes = item.marks_entry_deadline
	if not opens:
		return WINDOW_UNSCHEDULED, None, closes
	now = now_datetime()
	if now < get_datetime(opens):
		return WINDOW_NOT_OPEN, opens, closes
	if closes and now > get_datetime(closes):
		return WINDOW_CLOSED, opens, closes
	return WINDOW_OPEN, opens, closes


def _window_message(state, opens, closes):
	if state == WINDOW_UNSCHEDULED:
		return _("Submission deadlines are not set for this assessment yet.")
	if state == WINDOW_NOT_OPEN:
		return _("Marks entry opens after the submission deadline, on {0}.").format(_fmt(opens))
	if state == WINDOW_CLOSED:
		return _("The marks entry deadline ({0}) has passed.").format(_fmt(closes))
	return None


def _item_dict(item):
	state, opens, closes = _window(item)
	return {
		"name": item.name,
		"label": item.assessment_label,
		"component": item.component or "",
		"assessment_type": item.assessment_type or "",
		"assessment_date": _fmt(item.assessment_date),
		"first_deadline": _fmt(item.first_submission_deadline),
		"second_deadline": _fmt(item.second_submission_deadline),
		"marks_deadline": _fmt(item.marks_entry_deadline),
		"has_second": bool(item.second_submission_deadline),
		"remarks": item.remarks or "",
		"window": state,
		"window_message": _window_message(state, opens, closes),
		"opens_on": _fmt(opens),
	}


def _get_item(ctx, schedule_item):
	item = next((i for i in _schedule_items(ctx.offering) if i.name == schedule_item), None)
	if not item:
		frappe.throw(_("This assessment is not scheduled for your course offering."), frappe.PermissionError)
	return item


def _ctx_for(course_offering, item=None):
	"""faculty_grades context, on the schedule's exam plan when it has one.
	Student Course Marks is keyed by course + exam plan + student."""
	ctx = _context(course_offering, item.exam_plan if item and item.exam_plan else None)
	if item and not ctx.exam_plan:
		frappe.throw(_("No exam plan is set up for this course yet. Please contact the examination office."))
	return ctx


def _max_marks(ctx, item):
	"""Maximum marks from the Evaluation Schema: the config with the same
	component + assessment type, else the only config with that component."""
	if not item.component:
		return None
	configs = get_max_marks(ctx.evaluation_schema)
	exact = configs.get((item.component, item.assessment_type or ""))
	if exact:
		return exact
	same_component = [m for (comp, _atype), m in configs.items() if comp == item.component and m]
	return same_component[0] if len(same_component) == 1 else None


@frappe.whitelist()
def get_assessments(course_offering):
	"""Scheduled assessments of a course offering with their entry window."""
	ctx = _ctx_for(course_offering)
	return {
		"course_name": ctx.offering.course_name or ctx.course,
		"term": ctx.offering.term_name or "",
		"assessments": [_item_dict(i) for i in _schedule_items(ctx.offering)],
	}


@frappe.whitelist()
def get_marks_sheet(course_offering, schedule_item):
	"""Enrolled students with their 1st / 2nd submission marks for one assessment."""
	ctx = _ctx_for(course_offering)
	item = _get_item(ctx, schedule_item)
	ctx = _ctx_for(course_offering, item)
	return _sheet(ctx, item)


def _sheet(ctx, item):
	info = _item_dict(item)
	info["exam_plan"] = ctx.exam_plan
	info["max_marks"] = _max_marks(ctx, item)
	info["masked"] = bool(ctx.access.mask_student_info)
	info["locked"] = bool(ctx.access.locked)
	info["editable"] = info["window"] == WINDOW_OPEN and not info["locked"]
	if info["locked"]:
		info["window_message"] = _("Result entry is locked for this course.")

	students = _enrolled_students(ctx.offering.name)
	existing = {}
	if students:
		for r in frappe.db.sql(
			"""
			SELECT scm.student, scm.status,
			       IF(sam.first_marks_entered, sam.first_submission_marks, NULL) AS first_submission_marks,
			       IF(sam.second_marks_entered, sam.second_submission_marks, NULL) AS second_submission_marks,
			       sam.remarks, sam.entered_by, sam.entered_on
			FROM `tabStudent Course Marks` scm
			LEFT JOIN `tabStudent Assessment Marks` sam
			       ON sam.parent = scm.name AND sam.parenttype = 'Student Course Marks'
			      AND sam.schedule_item = %(item)s
			WHERE scm.course = %(course)s AND scm.exam_plan = %(plan)s AND scm.student IN %(students)s
			""",
			{"item": item.name, "course": ctx.course, "plan": ctx.exam_plan, "students": tuple(students)},
			as_dict=True,
		):
			existing[r.student] = r

	rows = []
	for s in _student_rows(students, info["masked"]):
		m = existing.get(s["student"]) or frappe._dict()
		status = m.status or STATUS_DRAFT
		s.update({
			"status": status,
			"row_editable": info["editable"] and status not in (STATUS_LOCKED, STATUS_SUBMITTED),
			"first": m.first_submission_marks,
			"second": m.second_submission_marks,
			"remarks": m.remarks or "",
			"entered_by": m.entered_by or "",
			"entered_on": _fmt(m.entered_on),
		})
		rows.append(s)

	return {"info": info, "students": rows}


@frappe.whitelist()
def save_marks(course_offering, schedule_item, rows):
	"""Save 1st / 2nd submission marks for several students.
	rows: [{student, first, second, remarks}] — only the changed students.
	Each student is saved on its own; failures are returned, not raised."""
	ctx = _ctx_for(course_offering)
	item = _get_item(ctx, schedule_item)
	ctx = _ctx_for(course_offering, item)
	_require_open(ctx, item)

	if isinstance(rows, str):
		rows = json.loads(rows)
	max_marks = _max_marks(ctx, item)
	has_second = bool(item.second_submission_deadline)

	saved, errors = [], []
	for row in rows or []:
		student = row.get("student")
		try:
			_assert_enrolled(ctx, student)
			first = parse_marks(row.get("first"), max_marks, _("1st Submission Marks"))
			second = parse_marks(row.get("second"), max_marks, _("2nd Submission Marks")) if has_second else None
			_save_student(ctx, item, student, first, second, (row.get("remarks") or "").strip(), max_marks)
			saved.append(student)
		except Exception as e:
			errors.append({"student": student, "error": frappe.utils.strip_html(str(e))})
			frappe.clear_messages()

	# No explicit commit: Frappe commits a successful POST request
	return {"saved": saved, "errors": errors}


def _require_open(ctx, item):
	if ctx.access.locked:
		frappe.throw(_("Result entry is locked for this course."), title=_("Not Editable"))
	state, opens, closes = _window(item)
	if state != WINDOW_OPEN:
		frappe.throw(_window_message(state, opens, closes), title=_("Marks Entry Closed"))


def _save_student(ctx, item, student, first, second, remarks, max_marks):
	scm = _scm(ctx, student, create=True)
	status = frappe.db.get_value("Student Course Marks", scm, "status") or STATUS_DRAFT
	if status == STATUS_LOCKED:
		frappe.throw(_("This student's marks are locked."))
	if status == STATUS_SUBMITTED:
		frappe.throw(_("This student's marks have been submitted. Ask the examination office to set them back to Draft."))

	# Float columns can't hold NULL, so a blank mark is stored as 0 with its
	# "entered" flag off — that keeps "not entered" apart from a real 0.
	values = {
		"assessment_label": item.assessment_label,
		"component": item.component,
		"assessment_type": item.assessment_type,
		"maximum_marks": flt(max_marks),
		"first_submission_marks": flt(first),
		"first_marks_entered": int(first is not None),
		"second_submission_marks": flt(second),
		"second_marks_entered": int(second is not None),
		"remarks": remarks,
		"assessment_schedule": item.assessment_schedule,
		"schedule_item": item.name,
		"entered_by": frappe.session.user,
		"entered_on": now_datetime(),
	}

	# Child rows are written directly: saving the parent would re-run its
	# total / grade calculation, which these marks don't feed (yet).
	existing = frappe.db.get_value(
		"Student Assessment Marks",
		{"parent": scm, "parenttype": "Student Course Marks", "schedule_item": item.name},
		"name",
	)
	if existing:
		frappe.db.set_value("Student Assessment Marks", existing, values)
		return

	idx = frappe.db.sql(
		"SELECT COALESCE(MAX(idx), 0) FROM `tabStudent Assessment Marks` WHERE parent = %s AND parenttype = 'Student Course Marks'",
		scm,
	)[0][0]
	child = frappe.get_doc({
		"doctype": "Student Assessment Marks",
		"parent": scm,
		"parenttype": "Student Course Marks",
		"parentfield": "assessment_marks",
		"idx": idx + 1,
		**values,
	})
	child.db_insert()


# ---------------------------------------------------------------------------
# Excel sheet: download / bulk upload
# ---------------------------------------------------------------------------
# Row 1 = headers, row 2 = hidden machine keys, column A = hidden student ID.
# Cell A2 also carries the schedule row ID, so a sheet can only be uploaded
# back into the assessment it was downloaded for.

SHEET_KEY = "key"


def _sheet_filename(ctx, item):
	name = f"{ctx.offering.course_name or ctx.course} - {item.assessment_label}"
	return re.sub(r'[\\/:*?"<>|]+', "-", name).strip() + ".xlsx"


@frappe.whitelist()
def download_sheet(course_offering, schedule_item):
	"""Excel sheet of the students listed for an assessment, with their
	current marks — fill it in and upload it with import_sheet."""
	import openpyxl
	from openpyxl.styles import Alignment, Font, PatternFill

	ctx = _ctx_for(course_offering)
	item = _get_item(ctx, schedule_item)
	ctx = _ctx_for(course_offering, item)
	data = _sheet(ctx, item)
	info, students = data["info"], data["students"]

	max_txt = f" (Max: {flt(info['max_marks']):g})" if info["max_marks"] else ""
	cols = [("Key", f"{SHEET_KEY}|{item.name}", 0), ("Registration ID", "reg", 18), ("Student Name", "name", 30),
		("1st Submission Marks" + max_txt, "first", 20)]
	if info["has_second"]:
		cols.append(("2nd Submission Marks" + max_txt, "second", 20))
	cols.append(("Remarks", "remarks", 40))

	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = (item.assessment_label or "Marks")[:31]
	for i, (header, key, width) in enumerate(cols, start=1):
		cell = ws.cell(row=1, column=i, value=header)
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = PatternFill("solid", fgColor="920C24")
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		ws.cell(row=2, column=i, value=key)
		if width:
			ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
	ws.row_dimensions[1].height = 32
	ws.row_dimensions[2].hidden = True
	ws.column_dimensions["A"].hidden = True
	ws.freeze_panes = "D3"

	field_of = {"first": "first", "second": "second", "remarks": "remarks", "reg": "reg_id", "name": "student_name"}
	for r_idx, s in enumerate(students, start=3):
		ws.cell(row=r_idx, column=1, value=s["student"])
		for c_idx, (_h, key, _w) in enumerate(cols[1:], start=2):
			ws.cell(row=r_idx, column=c_idx, value=s.get(field_of[key]))

	buf = io.BytesIO()
	wb.save(buf)
	frappe.response.filename = _sheet_filename(ctx, item)
	frappe.response.filecontent = buf.getvalue()
	frappe.response.type = "binary"


@frappe.whitelist()
def import_sheet(course_offering, schedule_item, filedata, filename="", apply=0):
	"""Check an uploaded sheet (base64) and, with apply=1, save it.

	Every row must be valid before anything is saved: the student is listed for
	this assessment and not Submitted / Locked, appears once, and marks are
	numbers within 0..max. A blank cell keeps the value already saved.
	Returns {ok, errors: [{row, student, column, message}], summary, applied}."""
	ctx = _ctx_for(course_offering)
	item = _get_item(ctx, schedule_item)
	ctx = _ctx_for(course_offering, item)
	_require_open(ctx, item)
	if not (filename or "").lower().endswith(".xlsx"):
		frappe.throw(_("Please upload the .xlsx sheet downloaded from this page."))

	import openpyxl

	try:
		raw = base64.b64decode((filedata or "").split(",")[-1])
		ws = openpyxl.load_workbook(io.BytesIO(raw), data_only=True).active
	except Exception:
		frappe.throw(_("The file could not be read as an Excel (.xlsx) workbook."))

	keys = [str(ws.cell(row=2, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
	headers = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
	if not keys or not keys[0].startswith(SHEET_KEY + "|"):
		frappe.throw(_("This is not the assessment marks sheet from this page. Download the sheet and fill that in."))
	if keys[0] != f"{SHEET_KEY}|{item.name}":
		frappe.throw(_("This sheet was downloaded for a different assessment. Download the sheet for {0} and fill that in.").format(
			frappe.bold(item.assessment_label)))

	col = {k: i for i, k in enumerate(keys)}
	has_second = bool(item.second_submission_deadline)
	max_marks = _max_marks(ctx, item)

	data = _sheet(ctx, item)
	listed = {s["student"]: s for s in data["students"]}
	reg_to_student = {s["reg_id"]: s["student"] for s in data["students"] if s["reg_id"] and s["reg_id"] != "—"}

	errors, pending, seen = [], {}, set()
	for row_no in range(3, ws.max_row + 1):
		row = [ws.cell(row=row_no, column=c + 1).value for c in range(len(keys))]
		if all(v is None or str(v).strip() == "" for v in row[1:]):
			continue
		sid = str(row[0] or "").strip()
		reg = str(row[col["reg"]] or "").strip() if "reg" in col else ""
		who = reg if reg and reg != "—" else (row[col["name"]] if "name" in col else sid) or sid
		if sid not in listed:
			sid = reg_to_student.get(reg, sid)
		s = listed.get(sid)
		if not s:
			errors.append({"row": row_no, "student": who, "column": "", "message": _("Student is not enrolled in this course offering.")})
			continue
		if sid in seen:
			errors.append({"row": row_no, "student": who, "column": "", "message": _("Student appears more than once in the file.")})
			continue
		seen.add(sid)
		if s["status"] in (STATUS_LOCKED, STATUS_SUBMITTED):
			errors.append({"row": row_no, "student": who, "column": "", "message": _("Marks are {0} and can't be changed.").format(s["status"])})
			continue

		# Blank cell = keep what is saved
		new = {"first": s["first"], "second": s["second"], "remarks": s["remarks"] or ""}
		changed = False
		for key, label in (("first", _("1st Submission Marks")), ("second", _("2nd Submission Marks"))):
			if key not in col or (key == "second" and not has_second):
				continue
			cell = row[col[key]]
			if cell is None or str(cell).strip() == "":
				continue
			try:
				new[key] = parse_marks(cell, max_marks, label)
				changed = True
			except frappe.ValidationError as e:
				errors.append({"row": row_no, "student": who, "column": headers[col[key]], "message": frappe.utils.strip_html(str(e))})
				frappe.clear_messages()
		if "remarks" in col and row[col["remarks"]] not in (None, ""):
			new["remarks"] = str(row[col["remarks"]]).strip()[:250]
			changed = True
		if changed:
			pending[sid] = new

	summary = {"students": len(pending)}
	if errors or not cint(apply):
		return {"ok": not errors, "errors": errors, "summary": summary, "applied": False}

	for sid, v in pending.items():
		_save_student(ctx, item, sid, v["first"], v["second"] if has_second else None, v["remarks"], max_marks)
	# No explicit commit: Frappe commits a successful POST request
	return {"ok": True, "errors": [], "summary": summary, "applied": True}
