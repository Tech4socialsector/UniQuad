# Copyright (c) 2026, TFSS and contributors
# For license information, please see license.txt

"""Year-to-year promotion.

Year-level promotion always runs against an Active Promotion Policy. Term-to-term
moves inside a year are done from Student Enrollment → Promote Students (Promotion
Run), which promotes without any policy check.

Every confirmed decision is kept as a Student Promotion record, which doubles as
the per-student Promotion Log (status, reasons, enrollment outcome, overrides).
"""

import frappe
from frappe import _
from frappe.utils import now_datetime, flt, cint, strip_html


PROMOTED_STATUSES     = ("Promoted", "Override - Promoted")
NOT_PROMOTED_STATUSES = ("Not Promoted", "Override - Not Promoted")
ALL_STATUSES          = PROMOTED_STATUSES + NOT_PROMOTED_STATUSES + ("Conditional",)
FEE_DUE_STATUSES      = ("Pending", "Partially Paid", "Overdue")

LOG_FIELDS = [
	"name", "student", "student_name", "batch_year", "programme",
	"current_year", "target_year", "promotion_status",
	"current_cgpa", "backlog_count", "attendance_percent",
	"shortage_course_count", "cf_fa_shortage_count",
	"cgpa_result", "backlog_result", "attendance_result",
	"shortage_course_result", "cf_result", "fee_due_result",
	"enrollment_status", "to_enrollment", "remarks",
	"manual_override", "override_reason", "processed_by", "processed_on",
	"stage", "published_by", "published_on", "notified_on",
]


# ── Filter option helpers ──────────────────────────────────────────────────────

@frappe.whitelist()
def get_programs():
	return frappe.db.get_all("Programme", fields=["name", "program_name"], order_by="name asc")


@frappe.whitelist()
def get_academic_years():
	return frappe.db.get_all(
		"Academic Year", fields=["name", "academic_year_name"],
		order_by="year_start_date desc"
	)


@frappe.whitelist()
def get_policies_for_filters(program, academic_year):
	"""Return matching Active policies for given program + academic_year."""
	if not program or not academic_year:
		return []
	return frappe.db.get_all(
		"Promotion Policy",
		filters={"program": program, "academic_year": academic_year, "status": "Active"},
		fields=["name", "title", "from_year", "to_year",
		        "enable_cgpa_check", "min_cgpa",
		        "enable_backlog_check", "max_backlogs_allowed",
		        "enable_attendance_check", "min_attendance_percent",
		        "enable_course_shortage_check", "max_shortage_courses",
		        "enable_cf_check", "max_cf_fa_shortage",
		        "block_on_fee_due",
		        "conditional_promotion_action", "auto_update_student_year"],
		order_by="from_year asc",
	)


# ── Core promotion engine ──────────────────────────────────────────────────────

def _evaluate_student(student_row, policy_dict):
	cgpa_result            = "Not Checked"
	backlog_result         = "Not Checked"
	attendance_result      = "Not Checked"
	shortage_course_result = "Not Checked"
	cf_result              = "Not Checked"
	failures               = 0

	if policy_dict.get("enable_cgpa_check"):
		cgpa = flt(student_row.get("current_cgpa") or 0)
		if cgpa >= flt(policy_dict.get("min_cgpa") or 0):
			cgpa_result = "Pass"
		else:
			cgpa_result = "Fail"
			failures += 1

	if policy_dict.get("enable_backlog_check"):
		backlogs = cint(student_row.get("backlog_count") or 0)
		if backlogs <= cint(policy_dict.get("max_backlogs_allowed") or 0):
			backlog_result = "Pass"
		else:
			backlog_result = "Fail"
			failures += 1

	if policy_dict.get("enable_attendance_check"):
		att = flt(student_row.get("attendance_percent") or 0)
		if att >= flt(policy_dict.get("min_attendance_percent") or 0):
			attendance_result = "Pass"
		else:
			attendance_result = "Fail"
			failures += 1

	# Criterion 3: No more than N courses with attendance shortage
	if policy_dict.get("enable_course_shortage_check"):
		shortage_count = cint(student_row.get("shortage_course_count") or 0)
		if shortage_count <= _max_shortage(policy_dict):
			shortage_course_result = "Pass"
		else:
			shortage_course_result = "Fail"
			failures += 1

	# Criterion 4: Carry-forward courses with FA applied but still not exam-eligible
	if policy_dict.get("enable_cf_check"):
		cf_shortage = cint(student_row.get("cf_fa_shortage_count") or 0)
		if cf_shortage <= cint(policy_dict.get("max_cf_fa_shortage") or 0):
			cf_result = "Pass"
		else:
			cf_result = "Fail"
			failures += 1

	total_checks = sum([
		bool(policy_dict.get("enable_cgpa_check")),
		bool(policy_dict.get("enable_backlog_check")),
		bool(policy_dict.get("enable_attendance_check")),
		bool(policy_dict.get("enable_course_shortage_check")),
		bool(policy_dict.get("enable_cf_check")),
	])

	if failures == 0:
		status = "Promoted"
	elif failures == total_checks:
		status = "Not Promoted"
	else:
		cond_action = policy_dict.get("conditional_promotion_action") or "Not Promoted"
		status = "Conditional" if "Conditional" in cond_action else "Not Promoted"

	return {
		"cgpa_result":            cgpa_result,
		"backlog_result":         backlog_result,
		"attendance_result":      attendance_result,
		"shortage_course_result": shortage_course_result,
		"cf_result":              cf_result,
		"promotion_status":       status,
	}


def _max_shortage(policy_dict):
	# 0 is a legitimate limit ("no shortage courses allowed"); only a blank
	# value falls back to the default of 2.
	val = policy_dict.get("max_shortage_courses")
	return 2 if val is None or val == "" else cint(val)


def _failure_summary(row, policy_dict, ev):
	"""Human-readable list of the criteria a student did not meet."""
	reasons = []
	if ev.get("cgpa_result") == "Fail":
		reasons.append(_("CGPA {0} below minimum {1}").format(
			flt(row.get("current_cgpa"), 2), flt(policy_dict.get("min_cgpa"), 2)))
	if ev.get("backlog_result") == "Fail":
		reasons.append(_("{0} backlog(s), max allowed {1}").format(
			cint(row.get("backlog_count")), cint(policy_dict.get("max_backlogs_allowed"))))
	if ev.get("attendance_result") == "Fail":
		reasons.append(_("Attendance {0}% below minimum {1}%").format(
			flt(row.get("attendance_percent"), 1), flt(policy_dict.get("min_attendance_percent"), 1)))
	if ev.get("shortage_course_result") == "Fail":
		reasons.append(_("{0} course(s) with attendance shortage, max allowed {1}").format(
			cint(row.get("shortage_course_count")), _max_shortage(policy_dict)))
	if ev.get("cf_result") == "Fail":
		reasons.append(_("{0} course(s) still short of attendance after FA/MFA (carry-forward), max allowed {1}").format(
			cint(row.get("cf_fa_shortage_count")), cint(policy_dict.get("max_cf_fa_shortage"))))
	if ev.get("fee_due_result") == "Fail":
		reasons.append(_("Outstanding fee due"))
	return "; ".join(reasons)


def _fee_due_students(student_names):
	"""Students with at least one unpaid Fee Demand."""
	if not student_names or not frappe.db.exists("DocType", "Fee Demand"):
		return set()
	return set(frappe.get_all(
		"Fee Demand",
		filters={"student": ["in", student_names], "status": ["in", FEE_DUE_STATUSES]},
		pluck="student",
		distinct=True,
	))


def _evaluate_year_promotion(row, policy_dict, has_fee_due):
	"""Full policy check for year-to-year promotion: the academic criteria
	plus the policy's fee-due block, with a readable reason string."""
	ev = _evaluate_student(row, policy_dict)
	ev["fee_due_result"] = "Not Checked"
	if policy_dict.get("block_on_fee_due"):
		if has_fee_due:
			ev["fee_due_result"]   = "Fail"
			ev["promotion_status"] = "Not Promoted"
		else:
			ev["fee_due_result"] = "Pass"
	ev["remarks"] = _failure_summary(row, policy_dict, ev)
	return ev


# A student's current Batch is the one on their active Student Enrollment;
# Student Master.batch is only the fallback (it is often blank or stale).
CURRENT_BATCH_SQL = """COALESCE(
	(SELECT se.batch FROM `tabStudent Enrollment` se
	 WHERE se.student = sm.name AND se.status = 'Enrolled' AND se.docstatus < 2
	   AND IFNULL(se.batch, '') != ''
	 ORDER BY se.enrollment_date DESC, se.creation DESC LIMIT 1),
	NULLIF(sm.batch, ''))"""

# Final grade = moderated/updated grade if present, else the original grade
# (same rule as Term Result). A grade fails when the course's Grading Schema
# flags it failed=1; without a schema assignment only "F" counts as a fail.
FINAL_GRADE_SQL = "COALESCE(NULLIF(scm.updated_grade, ''), NULLIF(scm.grade, ''))"
IMPROVEMENT_GRADES = ("C", "C+", "C⁺")


def _course_result_rows(student_names, academic_year, failed=True):
	"""Distinct (student, term, course) rows in the academic year that are
	failed (failed=True) or improvable C/C+ grades (failed=False)."""
	if not student_names:
		return []
	if failed:
		cond = f"(gsc.failed = 1 OR (gsc.name IS NULL AND {FINAL_GRADE_SQL} = 'F'))"
	else:
		cond = f"{FINAL_GRADE_SQL} IN %(imp)s AND IFNULL(gsc.failed, 0) = 0"
	return frappe.db.sql(
		f"""
		SELECT DISTINCT scm.student, ep.term AS term, scm.course,
		       COALESCE(NULLIF(c.course_name, ''), scm.course) AS course_name,
		       {FINAL_GRADE_SQL} AS final_grade, at2.term_start_date
		FROM `tabStudent Course Marks` scm
		INNER JOIN `tabExam Plan` ep ON ep.name = scm.exam_plan
		INNER JOIN `tabAcademic Term` at2 ON at2.name = ep.term
		LEFT JOIN `tabCourse` c ON c.name = scm.course
		LEFT JOIN `tabCourse Schema Assignment` csa
		       ON csa.exam_plan = scm.exam_plan AND csa.course = scm.course
		LEFT JOIN `tabGrading Schema Component` gsc
		       ON gsc.parent = csa.grade_schema AND gsc.parentfield = 'grades'
		      AND gsc.grade = {FINAL_GRADE_SQL}
		WHERE scm.student IN %(students)s
		  AND at2.academic_year = %(ay)s
		  AND COALESCE(scm.enrollment_status, '') NOT IN ('Dropped', 'Detained', 'Migrated')
		  AND {cond}
		ORDER BY at2.term_start_date, course_name
		""",
		{"students": student_names, "ay": academic_year, "imp": IMPROVEMENT_GRADES},
		as_dict=True,
	)


def _ordinal(n):
	n = cint(n)
	if 11 <= n % 100 <= 13:
		suffix = "th"
	else:
		suffix = {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
	return f"{n}{suffix}"


def _set_student_year(student, year):
	"""Write the study year in both formats Student Master uses."""
	frappe.db.set_value("Student Master", student, {
		"current_year": str(cint(year)),
		"year_of_study": f"{_ordinal(year)} Year",
	}, update_modified=False)


def _sync_student_master_batch(student, batch):
	"""Point Student Master at the student's new Batch. db.set_value skips
	fetch_from, so the Batch-derived fields are copied explicitly."""
	b = frappe.db.get_value(
		"Batch", batch, ["academic_year", "academic_term", "program", "level_of_study"], as_dict=True
	)
	if not b:
		return
	values = {"batch": batch, "academic_year": b.academic_year, "academic_term": b.academic_term}
	if b.program:
		values["programme_of_study"] = b.program
	if b.level_of_study:
		values["level_of_study"] = b.level_of_study
	frappe.db.set_value("Student Master", student, values, update_modified=False)


def _year_variants(year):
	"""Student Master.current_year is stored both as '1' and 'Year 1'."""
	y = str(year).strip()
	return list({y, f"Year {y}"})


def _get_students_raw(program, academic_year, from_year):
	"""
	Fetch active students for the given program + academic_year + from_year
	along with their CGPA, backlog count, and attendance average.
	"""
	students = frappe.db.sql(
		f"""
		SELECT
			s.student, s.first_name, s.current_cgpa, s.cur_batch AS programme,
			s.batch_year, s.current_year, c.batch_name AS cohort_name
		FROM (
			SELECT
				sm.name                  AS student,
				sm.first_name            AS first_name,
				sm.current_cgpa          AS current_cgpa,
				sm.year_of_study         AS batch_year,
				sm.current_year          AS current_year,
				{CURRENT_BATCH_SQL}      AS cur_batch
			FROM `tabStudent Master` sm
			WHERE sm.current_year IN %(from_year)s
			  AND sm.student_status = 'Active'
		) s
		INNER JOIN `tabBatch` c ON c.name = s.cur_batch
		WHERE c.program = %(program)s
		  AND c.academic_year = %(academic_year)s
		ORDER BY s.first_name
		""",
		{"program": program, "academic_year": academic_year, "from_year": _year_variants(from_year)},
		as_dict=True,
	)

	return _attach_metrics(students, academic_year)


def _attach_metrics(students, academic_year):
	"""Add backlog / attendance / shortage / carry-forward figures for the
	academic year to each student row (dicts with at least "student")."""
	if not students:
		return []

	student_names = [s["student"] for s in students]

	# Backlogs = distinct courses failed in this academic year's exams.
	# (Student Course Marks.status is Draft/Submitted/Locked — never "Fail" —
	# so failure is read from the final grade against the grading schema.)
	failed_courses = {}
	for r in _course_result_rows(student_names, academic_year, failed=True):
		failed_courses.setdefault(r.student, set()).add(r.course)
	backlog_map = {st: len(courses) for st, courses in failed_courses.items()}

	# Attendance average (overall %)
	att_rows = frappe.db.sql(
		"""
		SELECT student, AVG(attendance_percentage) AS avg_attendance
		FROM `tabAttendance Summary`
		WHERE student IN %(students)s
		  AND academic_year = %(academic_year)s
		GROUP BY student
		""",
		{"students": student_names, "academic_year": academic_year},
		as_dict=True,
	)
	att_map = {r["student"]: flt(r["avg_attendance"]) for r in att_rows}

	# Criterion 3: Count courses where student has attendance shortage
	# (attendance_percentage < minimum_required_percentage for the course)
	shortage_rows = frappe.db.sql(
		"""
		SELECT student, COUNT(*) AS shortage_count
		FROM `tabAttendance Summary`
		WHERE student IN %(students)s
		  AND academic_year = %(academic_year)s
		  AND attendance_percentage < minimum_required_percentage
		GROUP BY student
		""",
		{"students": student_names, "academic_year": academic_year},
		as_dict=True,
	)
	shortage_map = {r["student"]: cint(r["shortage_count"]) for r in shortage_rows}

	# Criterion 4: Carry-forward FA + shortage check
	# Counts courses where FA/MFA condonation was applied (total_fa_mfa_hours > 0)
	# but student is still not exam-eligible (eligible_for_exam = 0).
	cf_rows = frappe.db.sql(
		"""
		SELECT student, COUNT(*) AS cf_count
		FROM `tabAttendance Summary`
		WHERE student IN %(students)s
		  AND academic_year = %(academic_year)s
		  AND total_fa_mfa_hours > 0
		  AND eligible_for_exam = 0
		GROUP BY student
		""",
		{"students": student_names, "academic_year": academic_year},
		as_dict=True,
	)
	cf_map = {r["student"]: cint(r["cf_count"]) for r in cf_rows}

	for s in students:
		s["backlog_count"]         = backlog_map.get(s["student"], 0)
		s["attendance_percent"]    = att_map.get(s["student"], 0.0)
		s["shortage_course_count"] = shortage_map.get(s["student"], 0)
		s["cf_fa_shortage_count"]  = cf_map.get(s["student"], 0)
		s["student_name"]          = (s.get("first_name") or s.get("student_name") or "").strip()

	return students


def _next_academic_year(academic_year):
	"""The Academic Year that starts right after `academic_year`."""
	start = frappe.db.get_value("Academic Year", academic_year, "year_start_date")
	if not start:
		return None
	return frappe.db.get_value(
		"Academic Year", {"year_start_date": [">", start]}, "name", order_by="year_start_date asc"
	)


def _resolve_next_batch(current_cohort):
	"""Given a Student Enrollment's current Batch, find the Batch for the
	same program + entry cohort (section) one year_of_study ahead, in the
	next Academic Year by start date. Returns None if no such Batch exists
	yet (registrar hasn't created it) or the current batch has no section.

	NOTE: this only covers promotion that happens once per Academic Year.
	Term-to-term moves inside a year go through Promotion Run (Student
	Enrollment → Promote Students) instead.
	"""
	batch = frappe.db.get_value(
		"Batch", current_cohort,
		["program", "section", "term_year", "academic_year"], as_dict=True,
	)
	if not batch or not batch.section or batch.term_year is None:
		return None

	next_ay = _next_academic_year(batch.academic_year)
	if not next_ay:
		return None

	return frappe.db.get_value(
		"Batch",
		{
			"program": batch.program,
			"section": batch.section,
			"term_year": cint(batch.term_year) + 1,
			"academic_year": next_ay,
		},
		"name",
	)


def _promote_enrollment(student, to_year):
	"""Mark the student's current Student Enrollment Completed and create a
	new one in the resolved next Batch. Returns the new enrollment name, or
	raises if the next Batch doesn't exist yet."""
	current = frappe.db.get_value(
		"Student Enrollment",
		{"student": student, "status": "Enrolled", "docstatus": ["<", 2]},
		["name", "batch"],
		as_dict=True,
		order_by="enrollment_date desc",
	)
	if not current or not current.batch:
		frappe.throw(
			_("No active Student Enrollment found for {0} to promote from").format(student)
		)

	next_batch = _resolve_next_batch(current.batch)
	if not next_batch:
		frappe.throw(
			_(
				"Cannot promote {0}: no Batch found for year {1} following {2}. "
				"Create the target Batch first."
			).format(student, to_year, current.batch)
		)

	existing = frappe.db.exists(
		"Student Enrollment",
		{"student": student, "batch": next_batch, "docstatus": ["<", 2]},
	)
	if existing:
		_sync_student_master_batch(student, next_batch)
		return existing

	# Order matters: completing the old enrollment flips Student Master to
	# "Graduated"; inserting the new Enrolled one sets it back to "Active".
	old_doc = frappe.get_doc("Student Enrollment", current.name)
	old_doc.status = "Completed"
	old_doc.save(ignore_permissions=True)

	new_doc = frappe.new_doc("Student Enrollment")
	new_doc.student = student
	new_doc.batch = next_batch
	new_doc.status = "Enrolled"
	new_doc.enrollment_date = frappe.db.get_value("Batch", next_batch, "start_date") or frappe.utils.today()
	new_doc.insert(ignore_permissions=True)
	_sync_student_master_batch(student, next_batch)
	return new_doc.name


def _apply_year_promotion(student, to_year):
	"""Move the student to `to_year` and into the next year's Batch.

	Year, Batch and enrollment move together inside a savepoint: if the
	enrollment can't be created (e.g. target Batch not created yet) nothing
	changes — the student stays in their current year and Batch, the log row
	shows "Enrollment Failed", and Retry Enrollment completes the move later.
	"""
	save_point = "pm_enroll_" + frappe.generate_hash(length=10)
	msg_count = len(frappe.local.message_log or [])
	frappe.db.savepoint(save_point)
	try:
		name = _promote_enrollment(student, to_year)
		_set_student_year(student, to_year)
		frappe.db.release_savepoint(save_point)
		return {"status": "Enrolled", "to_enrollment": name, "error": None}
	except Exception as e:
		frappe.db.rollback(save_point=save_point)
		# frappe.throw() queues a popup even when caught; drop it — the
		# error is recorded on the log row instead.
		if frappe.local.message_log:
			frappe.local.message_log = frappe.local.message_log[:msg_count]
		error = strip_html(str(e)) or e.__class__.__name__
		frappe.log_error(
			title="Student Promotion: enrollment creation failed",
			message=f"Student {student}, target year {to_year}: {error}\n\n{frappe.get_traceback()}",
		)
		return {"status": "Failed", "to_enrollment": None, "error": error}


def _get_policy(policy_name, program=None, academic_year=None):
	if not policy_name:
		frappe.throw(_("Select a Promotion Policy. Year-to-year promotion always runs against a policy."))
	policy = frappe.get_doc("Promotion Policy", policy_name)
	if policy.status != "Active":
		frappe.throw(_("Promotion Policy {0} is {1}. Only Active policies can be used.").format(
			policy.name, policy.status))
	if program and policy.program != program:
		frappe.throw(_("Promotion Policy {0} belongs to Programme {1}, not {2}.").format(
			policy.name, policy.program, program))
	if academic_year and policy.academic_year != academic_year:
		frappe.throw(_("Promotion Policy {0} belongs to Academic Year {1}, not {2}.").format(
			policy.name, policy.academic_year, academic_year))
	return policy


def _check_policy_years(policy, from_year, to_year):
	"""A policy written with year levels (1–10) only applies to that step.
	Policies holding calendar years (e.g. 2026 → 2027) can't be matched and
	are allowed as-is (the page warns about them)."""
	if 1 <= cint(policy.from_year) <= 10 and (
		cint(from_year) != cint(policy.from_year) or (cint(to_year) and cint(to_year) != cint(policy.to_year))
	):
		frappe.throw(_("Promotion Policy {0} is for Year {1} → {2}, not Year {3} → {4}.").format(
			policy.name, policy.from_year, policy.to_year, cint(from_year), cint(to_year) or "?"))


def _validate_years(from_year, to_year):
	from_year, to_year = cint(from_year), cint(to_year)
	if from_year < 1:
		frappe.throw(_("Enter a valid Current Year (e.g. 1)."))
	if to_year <= from_year:
		frappe.throw(_("Target Year must be greater than Current Year."))
	return from_year, to_year


def _fill_missing_remarks(records, policy_name):
	"""Rows saved before reasons were recorded have no remarks — rebuild them
	from the stored check results and the policy's thresholds."""
	if not any(not r.get("remarks") and r.get("promotion_status") not in PROMOTED_STATUSES for r in records):
		return records
	policy_dict = frappe.get_doc("Promotion Policy", policy_name).as_dict()
	for r in records:
		if r.get("remarks") or r.get("promotion_status") in PROMOTED_STATUSES or r.get("manual_override"):
			continue
		r["remarks"] = _failure_summary(r, policy_dict, r) or _(
			"Did not meet the policy criteria (run before reasons were recorded — re-run to refresh)."
		)
	return records


def _count_statuses(rows):
	c = {"total": len(rows), "promoted": 0, "not_promoted": 0, "conditional": 0,
	     "enrolled": 0, "enrollment_failed": 0, "pending": 0}
	for r in rows:
		st = r.get("promotion_status")
		if st in PROMOTED_STATUSES:
			c["promoted"] += 1
		elif st in NOT_PROMOTED_STATUSES:
			c["not_promoted"] += 1
		else:
			c["conditional"] += 1
		if r.get("enrollment_status") == "Enrolled":
			c["enrolled"] += 1
		elif r.get("enrollment_status") == "Failed":
			c["enrollment_failed"] += 1
		elif r.get("enrollment_status") == "Pending":
			c["pending"] += 1
	return c


def get_student_progression(student):
	"""Year/term position and year-end promotion eligibility for one student,
	computed with the same engine as the Promotion Management page.

	Within a year, term-to-term moves are automatic (Student Enrollment →
	Promote Students). The Promotion Policy is only applied at year end.
	"""
	sm = frappe.db.get_value(
		"Student Master", student,
		["name", "first_name", "current_year", "current_cgpa", "student_status"], as_dict=True,
	)
	if not sm:
		return None
	batch = frappe.db.sql(f"SELECT {CURRENT_BATCH_SQL} FROM `tabStudent Master` sm WHERE sm.name = %s", student)
	batch = batch[0][0] if batch else None
	b = frappe.db.get_value(
		"Batch", batch, ["name", "program", "academic_year", "section", "term_year", "batch_name"], as_dict=True,
	) if batch else None

	import re as _re
	m = _re.search(r"\d+", str(sm.current_year or ""))
	year = int(m.group()) if m else 0

	out = {
		"year": year,
		"year_label": f"{_ordinal(year)} Year" if year else "",
		"batch": batch,
		"term_no": None, "terms_in_year": None, "is_last_term": None,
		"next_step": "", "eligibility": None, "last_decision": None,
	}

	# Term position inside the year: the Batch's term number among the
	# Batches of the same programme + section in that academic year.
	if b:
		out["term_no"] = cint(b.term_year) or None
		filters = {"program": b.program, "academic_year": b.academic_year}
		if b.section:
			filters["section"] = b.section
		term_years = [cint(t) for t in frappe.get_all("Batch", filters=filters, pluck="term_year") if cint(t)]
		out["terms_in_year"] = max(term_years) if term_years else None
		if out["term_no"] and out["terms_in_year"]:
			out["is_last_term"] = out["term_no"] >= out["terms_in_year"]
			if out["is_last_term"]:
				out["next_step"] = _("Year end — promotion to Year {0} is checked against the Promotion Policy.").format(year + 1) if year else ""
			else:
				out["next_step"] = _("Next: Term {0} of {1} — term-to-term promotion is automatic (no policy check).").format(
					out["term_no"] + 1, out["terms_in_year"])

	# Latest confirmed decision from the Promotion Log
	last = frappe.get_all(
		"Student Promotion", filters={"student": student, "stage": ["!=", "Superseded"]},
		fields=["name", "promotion_policy", "promotion_status", "current_year", "target_year",
		        "enrollment_status", "to_enrollment", "remarks", "processed_on", "manual_override",
		        "stage", "published_on"],
		order_by="processed_on desc, modified desc", limit=1,
	)
	out["last_decision"] = last[0] if last else None

	if not b or not year:
		return out

	# Policy for this student's year (Active, same programme + academic year).
	policies = frappe.get_all(
		"Promotion Policy",
		filters={"program": b.program, "academic_year": b.academic_year, "status": "Active"},
		fields=["name"], order_by="from_year asc, creation desc",
	)
	if not policies:
		return out
	docs = [frappe.get_doc("Promotion Policy", p.name) for p in policies]
	policy = next((d for d in docs if cint(d.from_year) == year), None)
	if not policy:
		# Some policies store calendar years (e.g. 2026) instead of year levels.
		policy = next((d for d in docs if cint(d.from_year) > 10), None)
	if not policy:
		return out

	row = _attach_metrics([{
		"student": student, "student_name": sm.first_name,
		"current_cgpa": sm.current_cgpa, "programme": batch,
	}], b.academic_year)[0]
	pd = policy.as_dict()
	ev = _evaluate_year_promotion(row, pd, bool(_fee_due_students([student])) if pd.get("block_on_fee_due") else False)

	criteria = []
	def add(enabled, label, value, required, result):
		if enabled:
			criteria.append({"label": label, "value": value, "required": required, "result": result})
	add(pd.enable_cgpa_check, _("CGPA"), f"{flt(row['current_cgpa']):.2f}", f"≥ {flt(pd.min_cgpa):.2f}", ev["cgpa_result"])
	add(pd.enable_backlog_check, _("Backlogs (failed courses)"), cint(row["backlog_count"]),
	    f"≤ {cint(pd.max_backlogs_allowed)}", ev["backlog_result"])
	add(pd.enable_attendance_check, _("Average attendance"), f"{flt(row['attendance_percent']):.1f}%",
	    f"≥ {flt(pd.min_attendance_percent):.1f}%", ev["attendance_result"])
	add(pd.enable_course_shortage_check, _("Courses with attendance shortage"), cint(row["shortage_course_count"]),
	    f"≤ {_max_shortage(pd)}", ev["shortage_course_result"])
	add(pd.enable_cf_check, _("Courses still short after FA/MFA (carry-forward)"), cint(row["cf_fa_shortage_count"]),
	    f"≤ {cint(pd.max_cf_fa_shortage)}", ev["cf_result"])
	add(pd.block_on_fee_due, _("Outstanding fee due"), _("Yes") if ev["fee_due_result"] == "Fail" else _("No"),
	    _("No"), ev["fee_due_result"])

	out["eligibility"] = {
		"policy": policy.name,
		"policy_title": policy.title,
		"academic_year": b.academic_year,
		"from_year": year,
		"to_year": year + 1,
		"status": ev["promotion_status"],
		"remarks": ev["remarks"],
		"criteria": criteria,
		"projected": not out["is_last_term"],
	}
	return out


# ── Public APIs ────────────────────────────────────────────────────────────────

@frappe.whitelist()
def fetch_students(program, academic_year, from_year, policy_name=None):
	"""Preview: evaluate every active student in `from_year` against the policy.
	Nothing is saved until confirm_promotion()."""
	frappe.has_permission("Student Promotion", "read", throw=True)
	policy = _get_policy(policy_name, program, academic_year)
	_check_policy_years(policy, from_year, 0)
	policy_dict = policy.as_dict()

	students = _get_students_raw(program, academic_year, from_year)
	if not students:
		return {"students": [], "counts": _count_statuses([])}

	fee_due = _fee_due_students([s["student"] for s in students]) if policy.block_on_fee_due else set()

	results = []
	for s in students:
		ev = _evaluate_year_promotion(s, policy_dict, s["student"] in fee_due)
		results.append({**s, **ev})

	return {"students": results, "counts": _count_statuses(results)}


@frappe.whitelist()
def confirm_promotion(program, academic_year, from_year, to_year, policy_name=None):
	"""
	Run Promotion: evaluate the students currently in `from_year` and save the
	decisions as **Draft**. Nothing is applied to students yet — no year change,
	no enrollment move, nothing visible on the portal, no notification.
	publish_promotion() applies and releases the drafts.

	Re-running replaces this step's earlier *Draft* rows for these students;
	Published rows stay until a new decision for the student is published.
	"""
	frappe.has_permission("Student Promotion", "create", throw=True)

	from_year, to_year = _validate_years(from_year, to_year)
	policy      = _get_policy(policy_name, program, academic_year)
	_check_policy_years(policy, from_year, to_year)
	policy_dict = policy.as_dict()
	students    = _get_students_raw(program, academic_year, from_year)

	if not students:
		frappe.throw(_("No active students found for the selected filters."))

	student_ids = [s["student"] for s in students]
	fee_due = _fee_due_students(student_ids) if policy.block_on_fee_due else set()

	old = frappe.get_all(
		"Student Promotion",
		filters={
			"promotion_policy": policy.name,
			"student": ["in", student_ids],
			"current_year": str(from_year),
			"target_year": str(to_year),
			"stage": "Draft",
		},
		pluck="name",
	)
	for name in old:
		frappe.delete_doc("Student Promotion", name, ignore_permissions=True, force=True)

	now = now_datetime()
	rows = []

	for s in students:
		ev     = _evaluate_year_promotion(s, policy_dict, s["student"] in fee_due)
		status = ev["promotion_status"]

		doc = frappe.new_doc("Student Promotion")
		doc.promotion_policy        = policy.name
		doc.student                 = s["student"]
		doc.student_name            = s["student_name"]
		doc.batch_year              = s.get("batch_year") or ""
		doc.programme               = s.get("programme") or ""
		doc.current_year            = str(from_year)
		doc.target_year             = str(to_year)
		doc.current_cgpa            = flt(s.get("current_cgpa") or 0)
		doc.backlog_count           = cint(s.get("backlog_count") or 0)
		doc.attendance_percent      = flt(s.get("attendance_percent") or 0)
		doc.shortage_course_count   = cint(s.get("shortage_course_count") or 0)
		doc.cf_fa_shortage_count    = cint(s.get("cf_fa_shortage_count") or 0)
		doc.cgpa_result             = ev["cgpa_result"]
		doc.backlog_result          = ev["backlog_result"]
		doc.attendance_result       = ev["attendance_result"]
		doc.shortage_course_result  = ev["shortage_course_result"]
		doc.cf_result               = ev["cf_result"]
		doc.fee_due_result          = ev["fee_due_result"]
		doc.promotion_status        = status
		doc.remarks                 = ev["remarks"]
		doc.stage                   = "Draft"
		doc.enrollment_status       = _pending_enrollment_status(status, policy)
		doc.processed_by            = frappe.session.user
		doc.processed_on            = now
		if status == "Promoted" and not policy.auto_update_student_year:
			doc.remarks = _("Auto-update of student year is off on the policy — year and enrollment will not change.")

		doc.insert(ignore_permissions=True)
		rows.append(doc.as_dict())

	frappe.db.commit()

	counts = _count_statuses(rows)
	return {
		"policy_name":  policy.name,
		"stage":        "Draft",
		"total":        counts["total"],
		"promoted":     counts["promoted"],
		"not_promoted": counts["not_promoted"],
		"conditional":  counts["conditional"],
	}


def _pending_enrollment_status(status, policy):
	"""Enrollment status a Draft row carries until it is published."""
	if status in PROMOTED_STATUSES and cint(policy.auto_update_student_year):
		return "Pending"
	return "Not Applicable"


def _log_records(policy_name, from_year=None, to_year=None, include_draft=True):
	"""Current log rows for a policy (optionally one year step): one row per
	student, the Draft winning over the Published row it will replace.
	Superseded rows are history and are left out."""
	filters = {"promotion_policy": policy_name, "stage": ["!=", "Superseded"]}
	if not include_draft:
		filters["stage"] = "Published"
	if cint(from_year):
		filters["current_year"] = str(cint(from_year))
	if cint(to_year):
		filters["target_year"] = str(cint(to_year))
	rows = frappe.get_all(
		"Student Promotion", filters=filters, fields=LOG_FIELDS,
		order_by="student_name asc", limit_page_length=0,
	)
	import datetime as _dt
	when = lambda r: r.processed_on or _dt.datetime.min
	best = {}
	for r in rows:
		key = (r.student, r.current_year, r.target_year)
		cur = best.get(key)
		if cur is None or (r.stage == "Draft" and cur.stage != "Draft") or (
			r.stage == cur.stage and when(r) > when(cur)
		):
			best[key] = r
	kept = set(id(v) for v in best.values())
	records = [r for r in rows if id(r) in kept]
	_fill_missing_remarks(records, policy_name)
	return records


def _stage_counts(records):
	return {
		"draft":     sum(1 for r in records if r.stage == "Draft"),
		"published": sum(1 for r in records if r.stage == "Published"),
	}


@frappe.whitelist()
def publish_promotion(policy_name, from_year, to_year, notify=0):
	"""Publish this step's Draft decisions:
	  - Promoted students move to the next year and are enrolled into next
	    year's Batch (atomically per student; failures stay in the current
	    year and can be retried from the log).
	  - Rows become Published; a Published row they replace is Superseded.
	  - Optionally emails each student their result (queued in background).
	"""
	frappe.has_permission("Student Promotion", "write", throw=True)
	from_year, to_year = _validate_years(from_year, to_year)
	policy = frappe.get_doc("Promotion Policy", policy_name)

	# Row locks: a second, simultaneous publish waits here and then finds none.
	drafts = [r[0] for r in frappe.db.sql(
		"""SELECT name FROM `tabStudent Promotion`
		   WHERE promotion_policy = %s AND current_year = %s AND target_year = %s AND stage = 'Draft'
		   ORDER BY student_name FOR UPDATE""",
		(policy.name, str(from_year), str(to_year)),
	)]
	if not drafts:
		frappe.throw(_("There are no Draft decisions to publish for this selection."))

	now = now_datetime()
	auto_update = cint(policy.auto_update_student_year)
	published, failures = [], []

	for name in drafts:
		doc = frappe.get_doc("Student Promotion", name)

		# The Published decision this Draft replaces becomes history.
		for old in frappe.get_all(
			"Student Promotion",
			filters={"promotion_policy": doc.promotion_policy, "student": doc.student,
			         "current_year": doc.current_year, "target_year": doc.target_year,
			         "stage": "Published", "name": ["!=", doc.name]},
			pluck="name",
		):
			frappe.db.set_value("Student Promotion", old, "stage", "Superseded", update_modified=False)

		if doc.promotion_status in PROMOTED_STATUSES and auto_update:
			current = str(frappe.db.get_value("Student Master", doc.student, "current_year") or "").strip()
			if current not in _year_variants(doc.current_year):
				doc.enrollment_status = "Failed"
				doc.remarks = _("Not moved: student is no longer in Year {0} (now {1}).").format(
					doc.current_year, current or "—")
				failures.append({"student": doc.student, "error": doc.remarks})
			else:
				outcome = _apply_year_promotion(doc.student, cint(doc.target_year))
				doc.enrollment_status = outcome["status"]
				doc.to_enrollment     = outcome["to_enrollment"]
				if outcome["error"]:
					doc.remarks = outcome["error"]
					failures.append({"student": doc.student, "error": outcome["error"]})
		elif doc.promotion_status not in PROMOTED_STATUSES:
			doc.enrollment_status = "Not Applicable"

		doc.stage        = "Published"
		doc.published_by = frappe.session.user
		doc.published_on = now
		doc.save(ignore_permissions=True)
		published.append(doc.as_dict())

	frappe.db.commit()

	queued, held = 0, 0
	if cint(notify):
		# A promoted student whose move failed is not told "promoted" yet.
		to_notify = []
		for r in published:
			if r.promotion_status in PROMOTED_STATUSES and r.enrollment_status == "Failed":
				held += 1
			elif frappe.db.get_value("Student Master", r.student, "email"):
				to_notify.append(r.name)
		queued = len(to_notify)
		if to_notify:
			frappe.enqueue(
				"slcm.slcm.page.promotion_management.promotion_management._notify_promotion_results",
				queue="short", timeout=1500, record_names=to_notify,
			)

	counts = _count_statuses(published)
	return {
		"published":           counts["total"],
		"promoted":            counts["promoted"],
		"not_promoted":        counts["not_promoted"],
		"conditional":         counts["conditional"],
		"enrolled":            counts["enrolled"],
		"enrollment_failures": failures,
		"notified_queued":     queued,
		"notify_held":         held,
	}


@frappe.whitelist()
def discard_draft(policy_name, from_year, to_year):
	"""Delete this step's Draft decisions (nothing was applied, so nothing to undo)."""
	frappe.has_permission("Student Promotion", "delete", throw=True)
	names = frappe.get_all(
		"Student Promotion",
		filters={"promotion_policy": policy_name, "current_year": str(cint(from_year)),
		         "target_year": str(cint(to_year)), "stage": "Draft"},
		pluck="name",
	)
	for name in names:
		frappe.delete_doc("Student Promotion", name, ignore_permissions=True, force=True)
	frappe.db.commit()
	return {"discarded": len(names)}


def _notify_promotion_results(record_names):
	"""Background job: email each student the published promotion result."""
	for name in record_names:
		try:
			rec = frappe.db.get_value(
				"Student Promotion", name,
				["name", "student", "promotion_status", "current_year", "target_year", "remarks",
				 "promotion_policy", "stage"], as_dict=True,
			)
			if not rec or rec.stage != "Published":
				continue
			if rec.promotion_status in PROMOTED_STATUSES and frappe.db.get_value(
				"Student Promotion", rec.name, "enrollment_status") == "Failed":
				continue  # not actually moved — don't announce a promotion
			sm = frappe.db.get_value("Student Master", rec.student, ["first_name", "email"], as_dict=True)
			if not sm or not sm.email:
				continue
			prog = frappe.db.get_value("Promotion Policy", rec.promotion_policy, "program") or ""
			prog = frappe.db.get_value("Programme", prog, "program_name") or prog
			to_ord, from_ord = _ordinal(rec.target_year), _ordinal(rec.current_year)
			if rec.promotion_status in PROMOTED_STATUSES:
				subject = _("Promotion result: promoted to {0} Year").format(to_ord)
				body = _("We are pleased to inform you that you have been promoted to <b>{0} Year</b> of {1}.").format(to_ord, prog)
			elif rec.promotion_status == "Conditional":
				subject = _("Promotion result: under review")
				body = _("Your promotion from {0} Year of {1} is under review. The office will contact you with the outcome.").format(from_ord, prog)
			else:
				reasons = "".join(f"<li>{frappe.utils.escape_html(x.strip())}</li>" for x in (rec.remarks or "").split(";") if x.strip())
				subject = _("Promotion result: not promoted")
				body = _("You have not met the promotion criteria for moving from {0} Year of {1}.").format(from_ord, prog) + (
					f"<ul>{reasons}</ul>" if reasons else "") + _("Please contact the office for the next steps.")
			frappe.sendmail(
				recipients=[sm.email],
				subject=subject,
				message=f"<p>{_('Dear {0},').format(frappe.utils.escape_html(sm.first_name or ''))}</p><p>{body}</p><p>{_('Regards,')}<br>{_('Office of the Registrar')}</p>",
				reference_doctype="Student Promotion",
				reference_name=rec.name,
			)
			frappe.db.set_value("Student Promotion", rec.name, "notified_on", now_datetime(), update_modified=False)
			frappe.db.commit()
		except Exception:
			frappe.db.rollback()
			frappe.log_error(title=f"Promotion result email failed: {name}", message=frappe.get_traceback())


@frappe.whitelist()
def get_promotion_log(policy_name, from_year=None, to_year=None):
	"""Saved per-student promotion log for a policy (optionally one year step)."""
	frappe.has_permission("Student Promotion", "read", throw=True)
	if not policy_name:
		return {"records": [], "counts": _count_statuses([]), "policy": None}

	records = _log_records(policy_name, from_year, to_year)
	policy = frappe.db.get_value(
		"Promotion Policy", policy_name,
		["name", "title", "program", "academic_year", "auto_update_student_year"], as_dict=True,
	)
	last = max((r.processed_on for r in records if r.processed_on), default=None)
	last_pub = max((r.published_on for r in records if r.published_on), default=None)
	return {
		"records": records,
		"counts": _count_statuses(records),
		"stages": _stage_counts(records),
		"policy": policy,
		"last_processed_on": last,
		"last_published_on": last_pub,
	}


@frappe.whitelist()
def get_saved_results_by_filters(program, academic_year, from_year, to_year, policy_name=None):
	"""Backward-compatible wrapper around get_promotion_log()."""
	frappe.has_permission("Student Promotion", "read", throw=True)
	if not policy_name:
		policy_name = frappe.db.get_value(
			"Promotion Policy",
			{"program": program, "academic_year": academic_year,
			 "from_year": cint(from_year), "to_year": cint(to_year)},
			"name",
		)
	log = get_promotion_log(policy_name, from_year, to_year) if policy_name else {"records": []}
	return {"records": log["records"], "policy_name": policy_name}


@frappe.whitelist()
def get_promotion_history(program, academic_year):
	"""One row per confirmed promotion (policy + year step) for the filters."""
	frappe.has_permission("Student Promotion", "read", throw=True)
	if not program or not academic_year:
		return []
	steps = frappe.db.sql(
		"""
		SELECT DISTINCT sp.promotion_policy, pp.title AS policy_title, sp.current_year, sp.target_year
		FROM `tabStudent Promotion` sp
		INNER JOIN `tabPromotion Policy` pp ON pp.name = sp.promotion_policy
		WHERE pp.program = %(program)s AND pp.academic_year = %(ay)s
		  AND IFNULL(sp.stage, '') != 'Superseded'
		""",
		{"program": program, "ay": academic_year}, as_dict=True,
	)
	out = []
	for st in steps:
		recs = _log_records(st.promotion_policy, st.current_year, st.target_year)
		if not recs:
			continue
		c, g = _count_statuses(recs), _stage_counts(recs)
		out.append({
			**st, **c,
			"drafts": g["draft"], "published_count": g["published"],
			"last_processed_on": max((r.processed_on for r in recs if r.processed_on), default=None),
			"last_published_on": max((r.published_on for r in recs if r.published_on), default=None),
		})
	import datetime as _dt
	out.sort(key=lambda x: x["last_processed_on"] or _dt.datetime.min, reverse=True)
	return out


@frappe.whitelist()
def save_override(record_name, new_status, reason):
	"""Manual override from the log.

	- Override - Promoted: actually promotes the student (year + next-year
	  enrollment, when the policy has auto-update on).
	- Override - Not Promoted: allowed while no new enrollment exists; the
	  student's year is moved back if it had already been advanced.
	"""
	frappe.has_permission("Student Promotion", "write", throw=True)
	if new_status not in ("Override - Promoted", "Override - Not Promoted"):
		frappe.throw(_("Invalid override status {0}.").format(new_status))
	reason = (reason or "").strip()
	if not reason:
		frappe.throw(_("A reason is required for a manual override."))

	doc = frappe.get_doc("Student Promotion", record_name)
	if doc.stage == "Superseded":
		frappe.throw(_("This decision has been superseded by a later one and can't be changed."))
	auto_update = cint(frappe.db.get_value("Promotion Policy", doc.promotion_policy, "auto_update_student_year"))
	was_promoted = doc.promotion_status in PROMOTED_STATUSES
	error = None
	stamp = f"[{frappe.utils.format_datetime(now_datetime())} · {frappe.session.user}]"

	if doc.stage == "Draft":
		# Draft: only the decision changes — it is applied on publish.
		if new_status == "Override - Promoted" and was_promoted:
			frappe.throw(_("{0} is already promoted.").format(doc.student_name or doc.student))
		if new_status == "Override - Not Promoted" and doc.promotion_status in NOT_PROMOTED_STATUSES:
			frappe.throw(_("{0} is already not promoted.").format(doc.student_name or doc.student))
		doc.promotion_status  = new_status
		doc.enrollment_status = "Pending" if (new_status == "Override - Promoted" and auto_update) else "Not Applicable"
		doc.to_enrollment     = None
		doc.manual_override   = 1
		doc.override_reason   = f"{stamp} {reason}"
		doc.remarks           = _("Manually overridden: {0}").format(reason)
		doc.save(ignore_permissions=True)
		frappe.db.commit()
		return {"ok": True, "enrollment_status": doc.enrollment_status, "error": None, "stage": "Draft"}

	if new_status == "Override - Promoted":
		if was_promoted:
			frappe.throw(_("{0} is already promoted.").format(doc.student_name or doc.student))
		if auto_update:
			outcome = _apply_year_promotion(doc.student, cint(doc.target_year))
			doc.enrollment_status = outcome["status"]
			doc.to_enrollment     = outcome["to_enrollment"]
			error = outcome["error"]
		else:
			doc.enrollment_status = "Not Applicable"
	else:
		if doc.promotion_status in NOT_PROMOTED_STATUSES:
			frappe.throw(_("{0} is already not promoted.").format(doc.student_name or doc.student))
		if doc.enrollment_status == "Enrolled" and doc.to_enrollment:
			frappe.throw(_(
				"{0} is already enrolled in {1} for the next year. Revert that enrollment "
				"manually before marking the student Not Promoted."
			).format(doc.student_name or doc.student, doc.to_enrollment))
		if was_promoted and auto_update:
			current = frappe.db.get_value("Student Master", doc.student, "current_year")
			if str(current or "").strip() in _year_variants(doc.target_year):
				_set_student_year(doc.student, doc.current_year)
		doc.enrollment_status = "Not Applicable"
		doc.to_enrollment     = None

	doc.promotion_status = new_status
	doc.manual_override  = 1
	doc.override_reason  = f"{stamp} {reason}"
	doc.remarks          = error or (_("Manually overridden: {0}").format(reason))
	doc.save(ignore_permissions=True)
	frappe.db.commit()
	return {"ok": True, "enrollment_status": doc.enrollment_status, "error": error, "stage": doc.stage}


def _recheck_record(doc, policy):
	"""Current promotion status of a log row's student under its policy."""
	academic_year = policy.academic_year
	cgpa = frappe.db.get_value("Student Master", doc.student, "current_cgpa")
	row = _attach_metrics([{"student": doc.student, "student_name": doc.student_name,
	                        "current_cgpa": cgpa}], academic_year)[0]
	pd = policy.as_dict()
	fee_due = bool(_fee_due_students([doc.student])) if pd.get("block_on_fee_due") else False
	return _evaluate_year_promotion(row, pd, fee_due)["promotion_status"]


@frappe.whitelist()
def retry_enrollment(record_names):
	"""Retry creating the next-year enrollment for promoted students whose
	enrollment failed (e.g. after the target Batch has been created)."""
	frappe.has_permission("Student Promotion", "write", throw=True)
	if isinstance(record_names, str):
		record_names = frappe.parse_json(record_names) if record_names.startswith("[") else [record_names]

	result = {"enrolled": [], "failed": []}
	for name in record_names or []:
		doc = frappe.get_doc("Student Promotion", name)
		if doc.promotion_status not in PROMOTED_STATUSES or doc.enrollment_status == "Enrolled":
			continue
		if doc.stage != "Published":
			result["failed"].append({"student": doc.student, "error": _("Not published yet — publish the draft first.")})
			continue
		policy = frappe.get_doc("Promotion Policy", doc.promotion_policy)
		if not cint(policy.auto_update_student_year):
			result["failed"].append({"student": doc.student, "error": _("Auto-update is off on the policy.")})
			continue
		# Policy-based "Promoted" rows are re-checked against current data before
		# moving the student (manual overrides are honoured as they are).
		if doc.promotion_status == "Promoted" and not doc.manual_override:
			still = _recheck_record(doc, policy)
			if still != "Promoted":
				result["failed"].append({"student": doc.student, "error": _(
					"No longer meets the policy ({0}). Run Promotion again to re-evaluate, or use Promote Anyway."
				).format(still)})
				continue
		outcome = _apply_year_promotion(doc.student, cint(doc.target_year))
		doc.enrollment_status = outcome["status"]
		doc.to_enrollment     = outcome["to_enrollment"]
		doc.remarks           = outcome["error"] or _("Enrolled on retry by {0}").format(frappe.session.user)
		doc.save(ignore_permissions=True)
		if outcome["error"]:
			result["failed"].append({"student": doc.student, "error": outcome["error"]})
		else:
			result["enrolled"].append(doc.student)
	frappe.db.commit()
	return result


@frappe.whitelist()
def download_promotion_list(policy_name, list_type, from_year=None, to_year=None):
	"""
	Download Excel for:
	  list_type = 'promoted' | 'not_promoted' | 'conditional' | 'enrollment_failed' | 'all'
	"""
	frappe.has_permission("Student Promotion", "read", throw=True)
	import io
	try:
		import openpyxl
		from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
		from openpyxl.utils import get_column_letter
	except ImportError:
		frappe.throw("openpyxl is not installed. Run: bench pip install openpyxl")

	status_map = {
		"promoted":          list(PROMOTED_STATUSES),
		"not_promoted":      list(NOT_PROMOTED_STATUSES),
		"conditional":       ["Conditional"],
		"enrollment_failed": list(PROMOTED_STATUSES),
		"all":               list(ALL_STATUSES),
	}
	label_map = {
		"promoted":          "Promoted List",
		"not_promoted":      "Not Promoted List",
		"conditional":       "Conditional List",
		"enrollment_failed": "Enrollment Failed List",
		"all":               "All Students",
	}
	if list_type not in status_map:
		frappe.throw(_("Unknown list type {0}.").format(list_type))

	policy  = frappe.get_doc("Promotion Policy", policy_name)
	wanted  = set(status_map[list_type])
	records = [
		r for r in _log_records(policy_name, from_year, to_year)
		if r.promotion_status in wanted and (list_type != "enrollment_failed" or r.enrollment_status == "Failed")
	]
	records.sort(key=lambda r: (r.promotion_status or "", (r.student_name or "").lower()))
	has_draft = any(r.stage == "Draft" for r in records)

	yr_from = cint(from_year) or policy.from_year
	yr_to   = cint(to_year) or policy.to_year

	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = label_map[list_type][:31]

	hdr_fill = PatternFill("solid", fgColor="920C24")
	hdr_font = Font(bold=True, color="FFFFFF", size=11)
	ctr      = Alignment(horizontal="center", vertical="center")
	wrap     = Alignment(vertical="top", wrap_text=True)
	bdr      = Border(
		left=Side(style="thin"), right=Side(style="thin"),
		top=Side(style="thin"),  bottom=Side(style="thin"),
	)
	row_colors = {
		"Promoted":                "D1FAE5",
		"Override - Promoted":     "A7F3D0",
		"Not Promoted":            "FEE2E2",
		"Override - Not Promoted": "FECACA",
		"Conditional":             "FEF3C7",
	}

	headers = ["#", "Stage", "Student ID", "Student Name", "Batch", "Current Year", "Target Year",
	           "CGPA", "Backlogs", "Attendance %", "Short Courses", "Short after FA/MFA",
	           "CGPA Check", "Backlog Check", "Attendance Check", "Shortage Check", "FA/MFA (CF) Check",
	           "Fee Due Check", "Promotion Status", "Enrollment Status", "New Enrollment",
	           "Remarks", "Override Reason", "Processed By", "Processed On"]
	col_widths = [5, 11, 18, 28, 14, 13, 13, 9, 10, 13, 16, 16, 13, 14, 16, 15, 12,
	              14, 22, 17, 20, 45, 35, 26, 20]
	last_col = get_column_letter(len(headers))

	ws.merge_cells(f"A1:{last_col}1")
	t = ws["A1"]
	t.value = (("DRAFT (not published) — " if has_draft else "") + f"{label_map[list_type]} — {policy.title}  "
	           f"({policy.program} | {policy.academic_year} | Year {yr_from} → Year {yr_to})")
	t.font      = Font(bold=True, size=13, color="6E0919")
	t.alignment = ctr
	ws.row_dimensions[1].height = 22

	counts = _count_statuses(records)
	ws.merge_cells(f"A2:{last_col}2")
	ws["A2"].value = (f"Promoted: {counts['promoted']}   |   "
	                  f"Not Promoted: {counts['not_promoted']}   |   "
	                  f"Conditional: {counts['conditional']}   |   "
	                  f"Enrolled: {counts['enrolled']}   |   "
	                  f"Enrollment Failed: {counts['enrollment_failed']}   |   "
	                  f"Total: {counts['total']}")
	ws["A2"].font = Font(italic=True, size=10, color="475569")
	ws.row_dimensions[2].height = 16

	for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
		cell           = ws.cell(row=3, column=ci, value=h)
		cell.fill      = hdr_fill
		cell.font      = hdr_font
		cell.alignment = ctr
		cell.border    = bdr
		ws.column_dimensions[get_column_letter(ci)].width = w
	ws.row_dimensions[3].height = 18

	for ri, rec in enumerate(records, 1):
		rn   = ri + 3
		rfil = PatternFill("solid", fgColor=row_colors.get(rec.promotion_status, "FFFFFF"))
		vals = [
			ri,
			rec.stage or "",
			rec.student,
			rec.student_name,
			rec.batch_year or "",
			rec.current_year or "",
			rec.target_year or "",
			round(flt(rec.current_cgpa), 2),
			cint(rec.backlog_count),
			str(round(flt(rec.attendance_percent), 1)) + "%",
			cint(rec.shortage_course_count),
			cint(rec.cf_fa_shortage_count),
			rec.cgpa_result or "Not Checked",
			rec.backlog_result or "Not Checked",
			rec.attendance_result or "Not Checked",
			rec.shortage_course_result or "Not Checked",
			rec.cf_result or "Not Checked",
			rec.fee_due_result or "Not Checked",
			rec.promotion_status,
			rec.enrollment_status or "",
			rec.to_enrollment or "",
			rec.remarks or "",
			rec.override_reason or "",
			rec.processed_by or "",
			frappe.utils.format_datetime(rec.processed_on) if rec.processed_on else "",
		]
		for ci, v in enumerate(vals, 1):
			cell        = ws.cell(row=rn, column=ci, value=v)
			cell.fill   = rfil
			cell.border = bdr
			cell.alignment = ctr if ci == 1 else wrap

	ws.freeze_panes = "E4"

	output = io.BytesIO()
	wb.save(output)
	output.seek(0)

	frappe.response.filename    = f"{list_type}_{policy_name}_Y{yr_from}-Y{yr_to}.xlsx"
	frappe.response.filecontent = output.read()
	frappe.response.type        = "download"


@frappe.whitelist()
def download_formatted_promotion_list(program, academic_year, university_name=None, include_draft=0):
	"""
	Official (NLS-style) promotion list for a Programme + Academic Year:
	  - One sheet per year level, taken from the confirmed promotion log
	  - Sections: Promoted · Conditional (pending review) · Re-admitted
	  - Term-wise failed (F) / attendance-shortage (AS) courses as columns
	  - C / C+ improvement courses and the reason for re-admission
	Falls back to a plain student list when no promotion has been run yet.
	"""
	frappe.has_permission("Student Promotion", "read", throw=True)
	import io
	from collections import OrderedDict
	try:
		import openpyxl
		from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
		from openpyxl.utils import get_column_letter as col_letter
	except ImportError:
		frappe.throw("openpyxl is not installed. Run: bench pip install openpyxl")

	prog_name = frappe.db.get_value("Programme", program, "program_name") or program
	univ      = (university_name or "").strip()

	# Terms of the academic year, in calendar order.
	terms = frappe.db.sql(
		"""SELECT name, term_name FROM `tabAcademic Term` WHERE academic_year = %s
		   ORDER BY term_start_date IS NULL, term_start_date, sequence, name""",
		academic_year, as_dict=True,
	)
	term_names  = [t.name for t in terms]
	term_labels = [t.term_name or t.name for t in terms]
	# Attendance Summary stores a free-text term name; map both spellings.
	term_lookup = {}
	for t in terms:
		term_lookup[t.name] = t.name
		if t.term_name:
			term_lookup.setdefault(t.term_name, t.name)

	# ── Confirmed decisions: latest per student + year ───────────────────────
	policy_names = frappe.get_all(
		"Promotion Policy",
		filters={"program": program, "academic_year": academic_year, "status": ["in", ["Active", "Draft"]]},
		pluck="name",
	)
	include_draft = cint(include_draft)
	stages = ["Published", "Draft"] if include_draft else ["Published"]
	records = frappe.get_all(
		"Student Promotion",
		filters={"promotion_policy": ["in", policy_names], "stage": ["in", stages]},
		fields=LOG_FIELDS + ["promotion_policy"],
		order_by="processed_on desc, modified desc",
		limit_page_length=0,
	) if policy_names else []
	# Drafts first (stable sort keeps newest-first) so a pending re-decision
	# wins over the published one it will replace.
	records.sort(key=lambda r: 0 if r.stage == "Draft" else 1)
	drafts_pending = bool(policy_names) and not include_draft and bool(frappe.db.exists(
		"Student Promotion", {"promotion_policy": ["in", policy_names], "stage": "Draft"}))
	has_draft_rows = any(r.stage == "Draft" for r in records)
	for pol in {r.promotion_policy for r in records}:
		_fill_missing_remarks([r for r in records if r.promotion_policy == pol], pol)

	# records are newest-first, so the first row per (student, year) is the
	# latest decision — a re-run under another policy supersedes older rows.
	by_year, seen = {}, set()
	for rec in records:
		year = cint(rec.current_year)
		if not year or (rec.student, year) in seen:
			continue
		seen.add((rec.student, year))
		by_year.setdefault(year, []).append(rec)
	by_year = OrderedDict(sorted(by_year.items()))

	# ── Fallback: nobody promoted yet → current student list per year ───────
	raw_mode = not by_year
	if raw_mode:
		raw = frappe.db.sql(
			f"""
			SELECT s.student, s.current_year, s.current_cgpa FROM (
				SELECT sm.name AS student, sm.current_year,
				       sm.current_cgpa AS current_cgpa, {CURRENT_BATCH_SQL} AS cur_batch
				FROM `tabStudent Master` sm WHERE sm.student_status = 'Active'
			) s
			INNER JOIN `tabBatch` c ON c.name = s.cur_batch
			WHERE c.program = %(program)s
			""",
			{"program": program}, as_dict=True,
		)
		import re as _re
		for s in raw:
			m = _re.search(r"\d+", str(s.current_year or ""))
			if not m:
				continue
			by_year.setdefault(int(m.group()), []).append(frappe._dict(
				student=s.student, current_cgpa=s.current_cgpa, promotion_status=None, remarks="",
			))
		by_year = OrderedDict(sorted(by_year.items()))
		if not by_year:
			frappe.throw(_("No active students found for <b>{0}</b>.").format(program))

	all_ids = list({r.student for rows in by_year.values() for r in rows})
	sm_rows = frappe.db.sql(
		"""SELECT name, email, TRIM(CONCAT_WS(' ', NULLIF(first_name, ''), NULLIF(middle_name, ''),
		          NULLIF(last_name, ''))) AS full_name
		   FROM `tabStudent Master` WHERE name IN %(ids)s""",
		{"ids": all_ids}, as_dict=True,
	)
	email_map = {r.name: r.email or "" for r in sm_rows}
	name_map  = {r.name: r.full_name or "" for r in sm_rows}

	# ── Course issues per student ────────────────────────────────────────────
	OTHER = "__other__"
	issues = {sid: {} for sid in all_ids}          # sid -> term -> [text]
	improve = {sid: [] for sid in all_ids}         # sid -> [text]
	fail_keys = set()
	for r in _course_result_rows(all_ids, academic_year, failed=True):
		tn = r.term if r.term in term_names else OTHER
		txt = f"{r.course_name} (F)"
		lst = issues[r.student].setdefault(tn, [])
		if txt not in lst:
			lst.append(txt)
		fail_keys.add((r.student, r.course))
	for r in frappe.db.sql(
		"""SELECT att.student, att.term_name, att.course,
		          COALESCE(NULLIF(c.course_name, ''), att.course) AS course_name
		   FROM `tabAttendance Summary` att
		   LEFT JOIN `tabCourse` c ON c.name = att.course
		   WHERE att.student IN %(ids)s AND att.academic_year = %(ay)s
		     AND att.attendance_percentage < att.minimum_required_percentage
		   ORDER BY course_name""",
		{"ids": all_ids, "ay": academic_year}, as_dict=True,
	):
		if (r.student, r.course) in fail_keys:
			continue  # already listed as failed
		tn = term_lookup.get(r.term_name) or OTHER
		txt = f"{r.course_name} (AS)"
		lst = issues[r.student].setdefault(tn, [])
		if txt not in lst:
			lst.append(txt)
	for r in _course_result_rows(all_ids, academic_year, failed=False):
		txt = f"{r.course_name} ({r.final_grade})"
		if txt not in improve[r.student]:
			improve[r.student].append(txt)

	# ── Styles ───────────────────────────────────────────────────────────────
	thin     = Side(style="thin", color="CBD5E1")
	bdr      = Border(left=thin, right=thin, top=thin, bottom=thin)
	ctr      = Alignment(horizontal="center", vertical="center", wrap_text=True)
	top_l    = Alignment(horizontal="left", vertical="top", wrap_text=True)
	hdr_fill = PatternFill("solid", fgColor="920C24")
	hdr_font = Font(bold=True, color="FFFFFF", size=10)
	SECTIONS = {
		"promoted":    (PatternFill("solid", fgColor="DCFCE7"), PatternFill("solid", fgColor="F0FDF4"), "166534"),
		"conditional": (PatternFill("solid", fgColor="FEF3C7"), PatternFill("solid", fgColor="FFFBEB"), "92400E"),
		"readmit":     (PatternFill("solid", fgColor="FEE2E2"), PatternFill("solid", fgColor="FFF7F7"), "991B1B"),
		"list":        (PatternFill("solid", fgColor="E2E8F0"), PatternFill("solid", fgColor="F8FAFC"), "0F172A"),
	}
	# Promoted and re-admitted students both continue in the *next* academic
	# year: prefer the year they were actually enrolled into.
	enrolled_ays = [
		ay for ay in (
			frappe.db.get_value("Student Enrollment", x.to_enrollment, "academic_year")
			for rows in by_year.values() for x in rows if x.get("to_enrollment")
		) if ay
	]
	next_ay = (max(set(enrolled_ays), key=enrolled_ays.count) if enrolled_ays
	           else _next_academic_year(academic_year))
	next_ay_txt = f" AY {next_ay}" if next_ay else ""
	first_t = term_labels[0] if term_labels else "Term 1"
	last_t  = term_labels[-1] if term_labels else "Last Term"

	def banner(ws, r, text, ncols, size=11, color="0F172A", bold=True, fill=None, height=18):
		ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
		c = ws.cell(row=r, column=1, value=text)
		c.font = Font(bold=bold, size=size, color=color)
		c.alignment = ctr
		if fill:
			c.fill = fill
		ws.row_dimensions[r].height = height
		return r + 1

	def write_section(ws, r, title, kind, recs, cols, ncols):
		sec_fill, row_fill, color = SECTIONS[kind]
		r = banner(ws, r, f"{title}  ({len(recs)} students)", ncols, color=color, fill=sec_fill)
		for ci, h in enumerate(cols, 1):
			cell = ws.cell(row=r, column=ci, value=h)
			cell.fill, cell.font, cell.alignment, cell.border = hdr_fill, hdr_font, ctr, bdr
		ws.row_dimensions[r].height = 32
		r += 1
		if not recs:
			ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
			c = ws.cell(row=r, column=1, value="— None —")
			c.alignment, c.font = ctr, Font(italic=True, color="64748B")
			ws.row_dimensions[r].height = 18
			return r + 2
		for si, rec in enumerate(recs, 1):
			sid = rec.student
			sname = name_map.get(sid) or rec.get("student_name") or sid
			if rec.get("stage") == "Draft":
				sname = f"{sname}  [DRAFT]"
			vals = [si, sid, sname,
			        email_map.get(sid, ""), round(flt(rec.current_cgpa), 2)]
			lines = 1
			for tn in term_cols:
				items = issues.get(sid, {}).get(tn, [])
				vals.append("\n".join(items))
				lines = max(lines, len(items))
			if "improve" in extras[kind]:
				vals.append("\n".join(improve.get(sid, [])))
				lines = max(lines, len(improve.get(sid, [])))
			if "reason" in extras[kind]:
				reason = (rec.remarks or "").replace("; ", "\n")
				if rec.promotion_status and rec.promotion_status not in ("Not Promoted", "Conditional"):
					reason = f"[{rec.promotion_status}]\n{reason}".strip()
				if rec.get("override_reason"):
					reason = f"{reason}\nOverride: {rec.override_reason}".strip()
				vals.append(reason)
				lines = max(lines, reason.count("\n") + 1)
			for ci, v in enumerate(vals, 1):
				cell = ws.cell(row=r, column=ci, value=v)
				cell.fill, cell.border = row_fill, bdr
				cell.alignment = ctr if ci in (1, 5) else top_l
			ws.row_dimensions[r].height = max(18, 15 * lines)
			r += 1
		return r + 1

	wb = openpyxl.Workbook()
	wb.remove(wb.active)

	for year, recs in by_year.items():
		from_ord = _ordinal(year)
		to_year  = max((cint(x.get("target_year")) for x in recs), default=0) or year + 1
		to_ord   = _ordinal(to_year)
		ws = wb.create_sheet(title=f"{from_ord} Year"[:31])

		sids = [x.student for x in recs]
		term_cols = list(term_names)
		term_hdrs = list(term_labels)
		if any(issues.get(s, {}).get(OTHER) for s in sids):
			term_cols.append(OTHER)
			term_hdrs.append("Other / unmapped term")
		base = ["Sl No", "Id No", "Student Name", "Email id", "CGPA"] + term_hdrs
		imp_hdr = f"C, C+ (Improvement Course {first_t} to {last_t}, if any)"
		extras = {
			"promoted":    ["improve"],
			"conditional": ["improve", "reason"],
			"readmit":     ["improve", "reason"],
			"list":        [],
		}
		cols_for = {
			"promoted":    base + [imp_hdr],
			"conditional": base + [imp_hdr, "Reason (pending review)"],
			"readmit":     base + [imp_hdr, "Reason for Re-admission"],
			"list":        base,
		}
		ncols = len(cols_for["readmit"]) if not raw_mode else len(base)

		r = 1
		if univ:
			r = banner(ws, r, univ, ncols, size=14, height=24)
		if raw_mode:
			r = banner(ws, r, f"Student List — {prog_name} — {from_ord} Year ({academic_year})", ncols, size=12)
			r = banner(ws, r, "(Promotion drafted but not yet published — showing currently enrolled students)"
			           if drafts_pending else "(Promotion not yet run — showing currently enrolled students)", ncols,
			           size=10, color="92400E", bold=False, height=15)
		else:
			r = banner(ws, r, f"Promotion List of {prog_name} — {from_ord} Year ({academic_year})", ncols, size=12)
			r = banner(ws, r, f"(Promoted to {to_ord} Year{next_ay_txt})", ncols,
			           size=10, color="374151", bold=False, height=15)
		if has_draft_rows and any(x.stage == "Draft" for x in recs):
			r = banner(ws, r, "DRAFT — includes decisions that are NOT yet published. Not an official list.",
			           ncols, size=11, color="FFFFFF", fill=PatternFill("solid", fgColor="B91C1C"), height=20)
		r = banner(ws, r, "(F) Failed   ·   (AS) Attendance shortage   ·   Improvement = C / C+ grades",
		           ncols, size=9, color="64748B", bold=False, height=14)
		freeze_row = r
		r += 1

		name_key = lambda x: (name_map.get(x.student) or x.get("student_name") or x.student).lower()
		if raw_mode:
			r = write_section(ws, r, f"Students in {from_ord} Year", "list", sorted(recs, key=name_key),
			                  cols_for["list"], ncols)
		else:
			promoted = sorted([x for x in recs if x.promotion_status in PROMOTED_STATUSES], key=name_key)
			cond     = sorted([x for x in recs if x.promotion_status == "Conditional"], key=name_key)
			readmit  = sorted([x for x in recs if x.promotion_status in NOT_PROMOTED_STATUSES], key=name_key)
			r = write_section(ws, r, f"Promoted to {to_ord} Year{next_ay_txt}", "promoted",
			                  promoted, cols_for["promoted"], ncols)
			if cond:
				r = write_section(ws, r, "Conditional — pending review", "conditional",
				                  cond, cols_for["conditional"], ncols)
			r = write_section(ws, r, f"Re-admitted to {from_ord} Year{next_ay_txt}", "readmit",
			                  readmit, cols_for["readmit"], ncols)

		widths = [7, 18, 28, 32, 9] + [24] * len(term_cols) + [34, 42]
		for ci, w in enumerate(widths[:ncols], 1):
			ws.column_dimensions[col_letter(ci)].width = w
		ws.freeze_panes = f"C{freeze_row}"
		ws.sheet_view.zoomScale = 90
		ws.page_setup.orientation = "landscape"
		ws.page_setup.fitToWidth = 1
		ws.sheet_properties.pageSetUpPr.fitToPage = True

	output = io.BytesIO()
	wb.save(output)
	output.seek(0)

	safe_prog = "".join(ch if ch.isalnum() else "_" for ch in program).strip("_")
	safe_ay   = academic_year.replace(" ", "").replace("-", "_")
	frappe.response.filename    = f"Promotion_List_{safe_prog}_{safe_ay}.xlsx"
	frappe.response.filecontent = output.read()
	frappe.response.type        = "download"
