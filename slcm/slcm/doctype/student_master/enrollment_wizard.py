# Copyright (c) 2026, TFSS and contributors
# For license information, please see license.txt

"""Server side of the Student Master list → Action → Enroll wizard.

Steps (see public/js/enrollment_wizard.js):
    1. Academic setup  - academic year (current by default), term, programme,
                         batch and year of study
    2. Students        - registration Completed, not yet enrolled in the programme
    3. Student ID      - a format such as M26#### (the #s are the running
                         number), assigned in alphabetical order
    4. Sections        - Mid-off / One-by-one (split by section capacities),
                         Manual, or Upload
    5. Summary         - enroll_students() saves everything
"""

import base64
import io
import re

import frappe
from frappe import _
from frappe.utils import cint, getdate, now_datetime, today

from slcm.slcm.doctype.student_master.student_master import (
    _append_section_update_log,
    _sync_enrollment_names,
)

ALLOWED_ROLES = ("System Manager", "slcm_Registrar")


def _check_permission():
    if frappe.session.user != "Administrator" and not set(ALLOWED_ROLES) & set(frappe.get_roles()):
        frappe.throw(_("You are not permitted to enroll students."), frappe.PermissionError)


def _full_name(s):
    return " ".join(filter(None, [s.get("first_name"), s.get("middle_name"), s.get("last_name")])).strip()


def _alpha_key(s):
    return (_full_name(s).lower(), s.get("name") or "")


def _parse_id_format(id_format):
    """'M26####' -> ('M26', 4, ''). Exactly one run of #s marks the number."""
    match = re.fullmatch(r"([^#]*)(#+)([^#]*)", (id_format or "").strip())
    if not match:
        frappe.throw(_("Student ID format must contain one run of # for the running number, e.g. M26####."))
    return match.group(1), len(match.group(2)), match.group(3)


def _existing_ids(prefix, width, suffix, exclude=()):
    """{number: student} of Student IDs already using this format."""
    pattern = re.compile(rf"{re.escape(prefix)}(\d{{{width}}}){re.escape(suffix)}")
    out = {}
    for r in frappe.get_all(
        "Student Master",
        filters={"registration_id": ["like", f"{prefix}%{suffix}"]},
        fields=["name", "registration_id"],
    ):
        match = pattern.fullmatch(r.registration_id or "")
        if match and r.name not in exclude:
            out[int(match.group(1))] = r.name
    return out


# ── Step 1 ────────────────────────────────────────────────────────────────────

@frappe.whitelist()
def get_wizard_defaults():
    """Current academic year and its current term."""
    academic_year = frappe.db.get_value("Academic Year", {"is_current_year": 1}, "name") \
        or frappe.db.get_value("Academic Year", {}, "name", order_by="year_start_date desc")
    return {"academic_year": academic_year, **get_year_defaults(academic_year)}


@frappe.whitelist()
def get_year_defaults(academic_year):
    """Default term (the one running today, else the first) and batch of an academic year."""
    if not academic_year:
        return {"academic_term": None, "batch": None}
    terms = frappe.get_all(
        "Academic Term",
        filters={"academic_year": academic_year, "status": "Active"},
        fields=["name", "term_start_date", "term_end_date"],
        order_by="term_no asc",
    )
    now = getdate(today())
    term = next((t.name for t in terms if getdate(t.term_start_date) <= now <= getdate(t.term_end_date)), None)
    return {
        "academic_term": term or (terms[0].name if terms else None),
        "batch": frappe.db.get_value("SLCM Batch", {"academic_year": academic_year}, "name", order_by="creation desc"),
    }


# ── Step 2 ────────────────────────────────────────────────────────────────────

@frappe.whitelist()
def get_enrollable_students(programme, academic_year):
    """Students with registration Completed in `academic_year` (Student
    Master's Academic Year), of this programme (or with no programme set
    yet), who are not enrolled in it yet and are not PACE students;
    alphabetical."""
    _check_permission()
    if not (programme and academic_year):
        return []
    enrolled = set(frappe.get_all("Student Enrollment", filters={"program": programme}, pluck="student"))
    students = frappe.get_all(
        "Student Master",
        filters={"registration_status": "Completed", "academic_year": academic_year},
        or_filters=[
            ["programme_of_study", "=", programme],
            ["programme_of_study", "is", "not set"],
        ],
        fields=["name", "first_name", "middle_name", "last_name", "registration_id",
                "gender", "batch", "section", "programme_of_study", "admission_type", "pace_programme"],
    )
    out = []
    for s in sorted(students, key=_alpha_key):
        if s.name in enrolled or s.admission_type == "PACE" or s.pace_programme:
            continue
        s.student_name = _full_name(s)
        out.append(s)
    return out


# ── Step 3 ────────────────────────────────────────────────────────────────────

def _id_conflicts(rows):
    """{student: other student} for students whose existing Student ID is also
    held by another Student Master (inside or outside the selection)."""
    ids = [r.registration_id for r in rows if r.registration_id]
    if not ids:
        return {}
    holders = {}
    for h in frappe.get_all("Student Master", filters={"registration_id": ["in", ids]},
                            fields=["name", "registration_id"]):
        holders.setdefault(h.registration_id, []).append(h.name)
    out = {}
    for r in rows:
        others = [n for n in holders.get(r.registration_id or "", []) if n != r.name]
        if others:
            out[r.name] = others[0]
    return out


@frappe.whitelist()
def preview_student_ids(students, id_format=None, start_from=None, renumber_existing=0,
                        numbering="Fill gaps", assign_new=None):
    """Order `students` alphabetically and work out each one's Student ID.

    An existing Student ID is never overwritten, except for students listed in
    `assign_new` or when `renumber_existing` is set. Everyone else without an
    ID gets the next number in `id_format` (e.g. M26####), skipping numbers
    already held: numbering "Fill gaps" starts from 1 (or `start_from`),
    "Continue after highest" right after the highest number in use.

    Each row has a status:
        new           - assigned now
        existing      - kept, matches the format
        other_format  - kept, but does not match the format (can be ticked
                        into assign_new)
        conflict      - kept ID is also held by another student; blocks enroll

    `id_format` may be left empty when nobody needs a new ID - e.g. every
    selected student already has one - and then nothing changes."""
    _check_permission()
    students = frappe.parse_json(students) if isinstance(students, str) else students
    assign_new = set(frappe.parse_json(assign_new) if isinstance(assign_new, str) else (assign_new or []))
    renumber_existing = cint(renumber_existing)

    rows = frappe.get_all(
        "Student Master",
        filters={"name": ["in", students or []]},
        fields=["name", "first_name", "middle_name", "last_name", "registration_id"],
    )
    rows.sort(key=_alpha_key)

    to_number = {r.name for r in rows if renumber_existing or not r.registration_id or r.name in assign_new}
    if to_number and not (id_format or "").strip():
        frappe.throw(_("{0} student(s) need a new Student ID: enter the Student ID format, e.g. M26####.").format(
            len(to_number)))

    pattern, taken, number, prefix, width, suffix = None, {}, 0, "", 0, ""
    if (id_format or "").strip():
        prefix, width, suffix = _parse_id_format(id_format)
        pattern = re.compile(rf"{re.escape(prefix)}\d{{{width}}}{re.escape(suffix)}")
        taken = _existing_ids(prefix, width, suffix, exclude=to_number)
        if cint(start_from):
            number = cint(start_from)
        elif numbering == "Continue after highest":
            number = max(taken, default=0) + 1
        else:
            number = 1

    conflicts = _id_conflicts([r for r in rows if r.name not in to_number])
    out = []
    for r in rows:
        if r.name in to_number:
            while number in taken:
                number += 1
            if len(str(number)) > width:
                frappe.throw(_("Student ID format {0} has room for {1} digits only: not enough free numbers for {2} students.").format(
                    id_format, width, len(to_number)))
            student_id, status = f"{prefix}{number:0{width}d}{suffix}", "new"
            taken[number] = r.name
            number += 1
        elif r.name in conflicts:
            student_id, status = r.registration_id, "conflict"
        elif pattern and not pattern.fullmatch(r.registration_id):
            student_id, status = r.registration_id, "other_format"
        else:
            student_id, status = r.registration_id, "existing"
        out.append({
            "student": r.name,
            "student_name": _full_name(r),
            "current_id": r.registration_id or "",
            "student_id": student_id,
            "status": status,
            "is_new": status == "new",
            "conflict_with": conflicts.get(r.name),
            "assign_new": r.name in assign_new,
        })

    count = lambda s: sum(1 for r in out if r["status"] == s)  # noqa: E731
    return {
        "rows": out,
        "summary": {
            "new": count("new"),
            "kept": len(out) - count("new"),
            "other_format": count("other_format"),
            "conflicts": count("conflict"),
        },
        "all_have_ids": not renumber_existing and all(r.registration_id for r in rows),
    }


# ── Step 4 (Upload) ───────────────────────────────────────────────────────────

@frappe.whitelist()
def download_section_template(rows):
    """xlsx of the wizard's students (alphabetical, with their Student ID) and
    an empty Section column to fill in."""
    _check_permission()
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    rows = frappe.parse_json(rows) if isinstance(rows, str) else rows
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sections"
    headers = ["Row Key", "Student ID", "Student Name", "Section"]
    ws.append(headers)
    for ci in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=ci)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3949AB")
    for r in rows or []:
        ws.append([r.get("student"), r.get("student_id") or "", r.get("student_name") or "", r.get("section") or ""])
    ws.column_dimensions["A"].hidden = True
    for col, width in zip("BCD", [16, 32, 14]):
        ws.column_dimensions[col].width = width

    output = io.BytesIO()
    wb.save(output)
    return {
        "filename": "enrollment_sections.xlsx",
        "content":  base64.b64encode(output.getvalue()).decode("utf-8"),
        "mime":     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }


@frappe.whitelist()
def read_section_upload(file_url):
    """{student: section} from a filled-in section template."""
    _check_permission()
    import openpyxl

    file_doc = frappe.get_doc("File", {"file_url": file_url})
    file_doc.check_permission("read")
    ws = openpyxl.load_workbook(file_doc.get_full_path(), data_only=True).active

    sections, unknown = {}, set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or len(row) < 4 or not row[3]:
            continue
        section = str(row[3]).strip()
        if not frappe.db.exists("Section", section):
            unknown.add(section)
        sections[str(row[0]).strip()] = section
    return {"sections": sections, "unknown": sorted(unknown)}


# ── Step 5 ────────────────────────────────────────────────────────────────────

@frappe.whitelist()
def enroll_students(data):
    """Enroll the wizard's students.

    data: academic_year, academic_term, programme, batch, year_no and rows
    [{student, student_id, section}]. For each student: Student Master gets
    the Student ID, section, batch, programme and academic year, then a
    Student Enrollment is created whose Academic Log starts at the chosen
    year and term. Each student is saved on its own: one failure does not
    stop the rest. Returns {success: [...], failed: [{student, reason}]}."""
    _check_permission()
    data = frappe._dict(frappe.parse_json(data) if isinstance(data, str) else data)
    for field, label in (("academic_year", _("Academic Year")), ("programme", _("Programme")),
                         ("batch", _("Batch"))):
        if not data.get(field):
            frappe.throw(_("{0} is required.").format(label))
    if data.academic_term and frappe.db.get_value("Academic Term", data.academic_term, "academic_year") != data.academic_year:
        frappe.throw(_("Academic Term {0} is not in {1}.").format(data.academic_term, data.academic_year))
    year_no = cint(data.year_no) or 1

    rows = data.rows or []
    ids = [r.get("student_id") for r in rows if r.get("student_id")]
    dupes = {i for i in ids if ids.count(i) > 1}
    if dupes:
        frappe.throw(_("Student ID used more than once: {0}").format(", ".join(sorted(dupes))))

    success, failed, renamed = [], [], []
    now = now_datetime()
    for r in rows:
        student = r.get("student")
        save_point = "enroll_wizard_" + frappe.generate_hash(length=8)
        frappe.db.savepoint(save_point)
        try:
            sm = frappe.db.get_value(
                "Student Master", student,
                ["name", "first_name", "middle_name", "last_name", "registration_status", "registration_id"],
                as_dict=True,
            )
            if not sm:
                raise frappe.ValidationError(_("Student Master not found."))
            if sm.registration_status != "Completed":
                raise frappe.ValidationError(_("Registration status is {0}, not Completed.").format(sm.registration_status))
            if frappe.db.exists("Student Enrollment", {"student": student, "program": data.programme}):
                raise frappe.ValidationError(_("Already enrolled in {0}.").format(data.programme))

            new_id = (r.get("student_id") or "").strip()
            if new_id:
                # New or kept, the ID must be this student's alone.
                holder = frappe.db.get_value("Student Master", {"registration_id": new_id, "name": ["!=", student]})
                if holder:
                    raise frappe.ValidationError(_("Student ID {0} is already used by {1}.").format(new_id, holder))
            section = (r.get("section") or "").strip()
            if section and not frappe.db.exists("Section", section):
                raise frappe.ValidationError(_("Section {0} does not exist.").format(section))

            values = {
                "batch": data.batch,
                "programme_of_study": data.programme,
                "academic_year": data.academic_year,
            }
            if new_id:
                values["registration_id"] = new_id
            if section:
                values["section"] = section
            frappe.db.set_value("Student Master", student, values)
            if section:
                _append_section_update_log(student, section, "Success",
                    _("Section set to '{0}' on enrollment.").format(section), now)

            enrollment = frappe.get_doc({
                "doctype":         "Student Enrollment",
                "student":         student,
                "student_name":    _full_name(sm),
                "program":         data.programme,
                "batch":           data.batch,
                "section":         section or None,
                "status":          "Enrolled",
                "enrollment_date": today(),
            })
            enrollment.add_year_log(data.academic_year, year_no, data.academic_term)
            # Roles were checked above; Registrars cannot create enrollments directly.
            enrollment.insert(ignore_permissions=True)

            if new_id and new_id != sm.registration_id:
                renamed.append(student)
            success.append({"student": student, "enrollment": enrollment.name})
        except Exception as e:
            frappe.db.rollback(save_point=save_point)
            frappe.clear_messages()
            if not isinstance(e, frappe.ValidationError):
                frappe.log_error(title=f"Enrollment wizard: {student}", message=frappe.get_traceback())
            failed.append({"student": student, "reason": frappe.utils.strip_html(str(e)) or e.__class__.__name__})
        else:
            frappe.db.release_savepoint(save_point)

    # Other enrollments of these students follow the new Student ID.
    _sync_enrollment_names(renamed)
    frappe.db.commit()
    return {"success": success, "failed": failed}
