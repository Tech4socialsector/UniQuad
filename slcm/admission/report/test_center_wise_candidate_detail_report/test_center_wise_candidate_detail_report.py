# Copyright (c) 2026, TFSS and contributors
# For license information, please see license.txt

import io
import json
import re
import zipfile
import frappe
from frappe import _
from frappe.utils.file_manager import save_file
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def execute(filters: dict | None = None):
	if filters is None:
		filters = {}

	columns = get_columns()
	data = get_data(filters)

	total_candidates = len(data)
	pwd_count = sum(1 for d in data if d.get("is_pwd") == "Yes")
	scribe_required_total = sum(1 for d in data if d.get("pwd_required_test") == "Yes")

	nlsiu_scribe_count = 0
	self_scribe_count = 0
	for d in data:
		if d.get("pwd_required_test") == "Yes":
			details = str(d.get("scribe_details", "")).lower()
			if "myself" in details or "self" in details:
				self_scribe_count += 1
			else:
				nlsiu_scribe_count += 1
	attended_count = sum(1 for d in data if d.get("entrance_test_status") == "Attended")
	absent_count = sum(1 for d in data if d.get("entrance_test_status") == "Absent")

	summary = [
		{
			"label": _("Total Candidates"),
			"value": total_candidates,
			"indicator": "Blue",
			"datatype": "Int",
		},
		{
			"label": _("PWD Candidates"),
			"value": pwd_count,
			"indicator": "Purple",
			"datatype": "Int",
		},
		{
			"label": _("Scribe Required"),
			"value": scribe_required_total,
			"indicator": "Orange",
			"datatype": "Int",
		},
		{
			"label": _("University Scribe Required"),
			"value": nlsiu_scribe_count,
			"indicator": "Yellow",
			"datatype": "Int",
		},
		{
			"label": _("Self Scribe Arranged"),
			"value": self_scribe_count,
			"indicator": "Green",
			"datatype": "Int",
		},
		{
			"label": _("Attended"),
			"value": attended_count,
			"indicator": "Green",
			"datatype": "Int",
		},
		{
			"label": _("Absent"),
			"value": absent_count,
			"indicator": "Red",
			"datatype": "Int",
		},
	]

	message = (
		_("Test Center-wise Candidate Detail Report.")
		if data
		else _("No candidates found matching the selected filters.")
	)

	chart = get_chart_data(data)

	return columns, data, message, chart, summary


def get_chart_data(data: list[dict]) -> dict | None:
	if not data:
		return None

	centre_map = {}
	for row in data:
		centre = row.get("test_centre") or _("Unknown")
		if centre not in centre_map:
			centre_map[centre] = {"univ_scribe": 0, "self_scribe": 0, "no_scribe": 0}

		req = row.get("pwd_required_test")
		details = str(row.get("scribe_details", "")).lower()

		if req == "Yes":
			if "myself" in details or "self" in details:
				centre_map[centre]["self_scribe"] += 1
			else:
				centre_map[centre]["univ_scribe"] += 1
		else:
			centre_map[centre]["no_scribe"] += 1

	sorted_centres = sorted(
		centre_map.keys(),
		key=lambda c: centre_map[c]["univ_scribe"] + centre_map[c]["self_scribe"] + centre_map[c]["no_scribe"],
		reverse=True,
	)[:15]

	labels = sorted_centres
	univ_vals = [centre_map[c]["univ_scribe"] for c in sorted_centres]
	self_vals = [centre_map[c]["self_scribe"] for c in sorted_centres]
	no_vals = [centre_map[c]["no_scribe"] for c in sorted_centres]

	return {
		"data": {
			"labels": labels,
			"datasets": [
				{
					"name": _("University Scribe Required"),
					"values": univ_vals,
				},
				{
					"name": _("Self Scribe Arranged"),
					"values": self_vals,
				},
				{
					"name": _("No Scribe Required"),
					"values": no_vals,
				},
			],
		},
		"type": "bar",
		"barOptions": {"stacked": True},
		"height": 300,
		"colors": ["#f39c12", "#27ae60", "#2980b9"],
	}


def get_columns() -> list[dict]:
	return [
		{
			"label": _("Sl No"),
			"fieldname": "sl_no",
			"fieldtype": "Int",
			"width": 70,
		},
		{
			"label": _("Test Centre"),
			"fieldname": "test_centre",
			"fieldtype": "Data",
			"width": 220,
		},
		{
			"label": _("Programme"),
			"fieldname": "program",
			"fieldtype": "Link",
			"options": "Programme",
			"width": 220,
		},
		{
			"label": _("Application Number"),
			"fieldname": "applicant",
			"fieldtype": "Link",
			"options": "Applicant",
			"width": 160,
		},
		{
			"label": _("Admit Card Number"),
			"fieldname": "admit_card_number",
			"fieldtype": "Data",
			"width": 200,
		},
		{
			"label": _("Student Name"),
			"fieldname": "candidate_name",
			"fieldtype": "Data",
			"width": 220,
		},
		{
			"label": _("Is PWD?"),
			"fieldname": "is_pwd",
			"fieldtype": "Data",
			"width": 90,
		},
		{
			"label": _("Scribe Required"),
			"fieldname": "pwd_required_test",
			"fieldtype": "Data",
			"width": 130,
		},
		{
			"label": _("Scribe Details"),
			"fieldname": "scribe_details",
			"fieldtype": "Data",
			"width": 280,
		},
		{
			"label": _("Attendance / Status"),
			"fieldname": "entrance_test_status",
			"fieldtype": "Data",
			"width": 140,
		},
	]


def get_data(filters: dict) -> list[dict]:
	user = frappe.session.user
	roles = frappe.get_roles(user)

	provider_filter = None
	if (
		"Entrance Test Provider" in roles
		and "System Manager" not in roles
		and "Entrance Test Admin" not in roles
		and user != "Administrator"
	):
		provider_name = frappe.db.get_value("Entrance Test Provider", {"user": user}, "name")
		if not provider_name:
			return []
		provider_filter = provider_name
	elif filters.get("entrance_test_provider"):
		provider_filter = filters.get("entrance_test_provider")

	conditions = ["sa.docstatus < 2"]
	values = {}

	# Only filter by PWD if explicitly requested
	if filters.get("is_pwd_only") or filters.get("pwd_only"):
		conditions.append("(sa.pwd = 1 OR app.pwd = 'Yes')")

	if filters.get("scribe_required_only"):
		conditions.append("(app.pwd_required_test = 'Yes' OR (app.scribe_allotment IS NOT NULL AND app.scribe_allotment != ''))")

	if filters.get("city"):
		conditions.append("(prov.city = %(city)s OR sa.center_name LIKE %(city_like)s OR sa.re_center_name LIKE %(city_like)s)")
		values["city"] = filters.get("city")
		values["city_like"] = f"%{filters.get('city')}%"

	if filters.get("is_international_applicant") or filters.get("show_international_applicant"):
		conditions.append(
			"(sa.is_international_applicant = 1 OR sa.entrance_test_provider = 'International Applicant' OR sa.entrance_test_provider IS NULL OR sa.entrance_test_provider = '')"
		)
	elif provider_filter:
		conditions.append(
			"(sa.entrance_test_provider = %(provider)s OR sa.re_entrance_test_provider = %(provider)s OR sa.center_name = %(provider)s OR sa.re_center_name = %(provider)s)"
		)
		values["provider"] = provider_filter

	if filters.get("academic_year"):
		conditions.append("sa.academic_year = %(academic_year)s")
		values["academic_year"] = filters.get("academic_year")

	if filters.get("admission_cycle"):
		conditions.append("sa.admission_cycle = %(admission_cycle)s")
		values["admission_cycle"] = filters.get("admission_cycle")

	if filters.get("campus"):
		conditions.append("sa.campus = %(campus)s")
		values["campus"] = filters.get("campus")

	if filters.get("program_level"):
		conditions.append("sa.program_level = %(program_level)s")
		values["program_level"] = filters.get("program_level")

	if filters.get("program"):
		conditions.append("sa.program = %(program)s")
		values["program"] = filters.get("program")

	if filters.get("allocation_date"):
		conditions.append("sa.allocation_date = %(allocation_date)s")
		values["allocation_date"] = filters.get("allocation_date")

	if filters.get("entrance_test_status"):
		conditions.append("sa.entrance_test_status = %(entrance_test_status)s")
		values["entrance_test_status"] = filters.get("entrance_test_status")

	where_clause = " AND ".join(conditions)

	sql_query = f"""
		SELECT
			sa.name as sa_name,
			sa.applicant,
			sa.candidate_name,
			sa.program,
			sa.academic_year,
			sa.admit_card_number,
			sa.center_name,
			sa.re_center_name,
			sa.entrance_test_provider,
			sa.re_entrance_test_provider,
			sa.is_rescheduled,
			sa.pwd as sa_pwd,
			sa.entrance_test_status,
			app.applicant_id,
			app.pwd as app_pwd,
			app.pwd_required_test,
			app.scribe_allotment
		FROM `tabEntrance Test Seat Allocation` sa
		LEFT JOIN `tabApplicant` app ON sa.applicant = app.name
		LEFT JOIN `tabEntrance Test Provider` prov ON prov.name = IFNULL(
			NULLIF(
				IF(
					sa.is_rescheduled = 1
					AND sa.re_entrance_test_provider IS NOT NULL
					AND sa.re_entrance_test_provider != '',
					sa.re_entrance_test_provider,
					sa.entrance_test_provider
				),
				''
			),
			NULL
		)
		WHERE {where_clause}
		ORDER BY
			COALESCE(
				NULLIF(IF(sa.is_rescheduled = 1 AND sa.re_center_name IS NOT NULL AND sa.re_center_name != '', sa.re_center_name, sa.center_name), ''),
				NULLIF(IF(sa.is_rescheduled = 1 AND sa.re_entrance_test_provider IS NOT NULL AND sa.re_entrance_test_provider != '', sa.re_entrance_test_provider, sa.entrance_test_provider), ''),
				'International Applicant'
			) ASC,
			sa.candidate_name ASC
	"""

	records = frappe.db.sql(sql_query, values, as_dict=True)

	data = []
	for idx, r in enumerate(records, start=1):
		# Test centre determination
		if r.get("is_rescheduled") == 1 and (r.get("re_center_name") or r.get("re_entrance_test_provider")):
			centre = r.get("re_center_name") or r.get("re_entrance_test_provider")
		else:
			centre = r.get("center_name") or r.get("entrance_test_provider") or "International Applicant"

		# PWD status determination
		is_pwd_val = "Yes" if (r.get("sa_pwd") == 1 or (r.get("app_pwd") or "").strip().lower() == "yes") else "No"

		# Scribe details determination
		req_scribe = (r.get("pwd_required_test") or "").strip()
		allotment = (r.get("scribe_allotment") or "").strip()

		if req_scribe.lower() == "yes" or allotment:
			scribe_req_disp = "Yes"
			if allotment:
				scribe_info = allotment
			else:
				scribe_info = "Yes (Arrangement Not Specified)"
		elif req_scribe.lower() == "no":
			scribe_req_disp = "No"
			scribe_info = "No Scribe Required"
		else:
			scribe_req_disp = "Not Specified" if is_pwd_val == "Yes" else "No"
			scribe_info = "-" if is_pwd_val == "No" else "Not Specified"

		# Admit Card Number & Application Number
		application_no = r.get("applicant") or r.get("applicant_id") or ""
		admit_card_no = r.get("admit_card_number") or (f"AC-{application_no}" if application_no else "")

		data.append({
			"sl_no": idx,
			"test_centre": centre,
			"program": r.get("program") or "",
			"academic_year": r.get("academic_year") or "",
			"applicant": application_no,
			"admit_card_number": admit_card_no,
			"candidate_name": r.get("candidate_name") or "",
			"is_pwd": is_pwd_val,
			"pwd_required_test": scribe_req_disp,
			"scribe_details": scribe_info,
			"entrance_test_status": r.get("entrance_test_status") or "Scheduled",
		})

	return data


def _safe_filename(name: str) -> str:
	clean = re.sub(r'[\\/*?:"<>|]', '_', str(name or "")).strip()
	return clean or "Test_Centre"


def _format_candidate_sheet(ws, title_text: str, rows: list[dict], include_centre: bool = False):
	ws.views.sheetView[0].showGridLines = True

	if include_centre:
		headers = [
			"Sl No", "Test Centre", "Programme", "Application Number",
			"Admit Card Number", "Student Name", "Is PWD?", "Scribe Required", "Attendance"
		]
		alignments = ["center", "left", "left", "center", "center", "left", "center", "center", "center"]
		widths = [10, 32, 28, 22, 24, 30, 12, 16, 14]
	else:
		headers = [
			"Sl No", "Programme", "Application Number",
			"Admit Card Number", "Student Name", "Is PWD?", "Scribe Required", "Attendance"
		]
		alignments = ["center", "left", "center", "center", "left", "center", "center", "center"]
		widths = [10, 28, 22, 24, 30, 12, 16, 14]

	last_col = len(headers)
	last_letter = get_column_letter(last_col)

	# Row 1: Merged Title
	ws.merge_cells(f"A1:{last_letter}1")
	t_cell = ws["A1"]
	t_cell.value = title_text
	t_cell.font = Font(name="Calibri", size=13, bold=True, color="111827")
	t_cell.alignment = Alignment(horizontal="center", vertical="center")
	ws.row_dimensions[1].height = 30

	# Row 2: Subtitle instruction
	ws.merge_cells(f"A2:{last_letter}2")
	s_cell = ws["A2"]
	s_cell.value = "Enter AB for absent students in the Attendance column"
	s_cell.font = Font(name="Calibri", size=10, italic=True, color="4B5563")
	s_cell.alignment = Alignment(horizontal="left", vertical="center")
	ws.row_dimensions[2].height = 20

	# Row 3: Blank separator
	ws.row_dimensions[3].height = 12

	# Row 4: Header
	ws.row_dimensions[4].height = 25
	hdr_border = Border(
		left=Side(style="thin", color="9CA3AF"),
		right=Side(style="thin", color="9CA3AF"),
		top=Side(style="thin", color="9CA3AF"),
		bottom=Side(style="thin", color="9CA3AF")
	)
	hdr_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

	for c_idx, h in enumerate(headers, start=1):
		c = ws.cell(row=4, column=c_idx, value=h)
		c.font = Font(name="Calibri", size=11, bold=True, color="111827")
		c.fill = hdr_fill
		c.alignment = Alignment(horizontal=alignments[c_idx - 1], vertical="center")
		c.border = hdr_border
		ws.column_dimensions[get_column_letter(c_idx)].width = widths[c_idx - 1]

	# Row 5+: Data rows
	data_border = Border(
		left=Side(style="thin", color="E5E7EB"),
		right=Side(style="thin", color="E5E7EB"),
		top=Side(style="thin", color="E5E7EB"),
		bottom=Side(style="thin", color="E5E7EB")
	)
	alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

	for row_num, r in enumerate(rows, start=5):
		ws.row_dimensions[row_num].height = 22
		sl = row_num - 4

		# Formatting values
		is_pwd_display = "Yes" if r.get("is_pwd") == "Yes" else ""
		scribe_display = "Yes" if r.get("pwd_required_test") == "Yes" else ""
		attendance_display = "AB" if r.get("entrance_test_status") == "Absent" else ""

		if include_centre:
			row_vals = [
				sl,
				r.get("test_centre") or "",
				r.get("program") or "",
				r.get("applicant") or "",
				r.get("admit_card_number") or "",
				r.get("candidate_name") or "",
				is_pwd_display,
				scribe_display,
				attendance_display,
			]
		else:
			row_vals = [
				sl,
				r.get("program") or "",
				r.get("applicant") or "",
				r.get("admit_card_number") or "",
				r.get("candidate_name") or "",
				is_pwd_display,
				scribe_display,
				attendance_display,
			]

		for c_idx, val in enumerate(row_vals, start=1):
			c = ws.cell(row=row_num, column=c_idx, value=val)
			c.font = Font(name="Calibri", size=10, color="111827")
			c.alignment = Alignment(horizontal=alignments[c_idx - 1], vertical="center")
			c.border = data_border
			if row_num % 2 == 0:
				c.fill = alt_fill


@frappe.whitelist()
def download_candidate_detail_excel(filters=None, download_type="overall"):
	if isinstance(filters, str):
		try:
			filters = json.loads(filters)
		except Exception:
			filters = {}
	if not filters:
		filters = {}

	data = get_data(filters)
	if not data:
		frappe.throw(_("No candidate records found matching the selected filters."))

	ts = frappe.utils.now_datetime().strftime("%Y%m%d_%H%M%S")

	# Resolve academic year if available
	ay = filters.get("academic_year")
	if not ay and data:
		record_ays = {d.get("academic_year") for d in data if d.get("academic_year")}
		if len(record_ays) == 1:
			ay = list(record_ays)[0]

	ay_part = ""
	if ay:
		cleaned_ay = re.sub(r'[^a-zA-Z0-9]+', '_', str(ay)).strip('_')
		if cleaned_ay:
			ay_part = f"_{cleaned_ay}"

	if download_type == "overall":
		# Single Excel file containing overall records
		wb = openpyxl.Workbook()
		ws = wb.active
		ws.title = "Candidate Attendance"

		# If filtered by a specific center, use that center name as header
		single_center = filters.get("entrance_test_provider")
		distinct_centres = {d.get("test_centre") for d in data}
		if single_center or len(distinct_centres) == 1:
			title_text = single_center or list(distinct_centres)[0]
			_format_candidate_sheet(ws, title_text, data, include_centre=False)
		else:
			title_text = "All Test Centres - Overall Candidate Attendance"
			_format_candidate_sheet(ws, title_text, data, include_centre=True)

		buf = io.BytesIO()
		wb.save(buf)
		buf.seek(0)

		filename = f"Candidate_Details_Overall{ay_part}_{ts}.xlsx"
		saved = save_file(
			filename,
			buf.getvalue(),
			"Report",
			"Test Center-wise Candidate Detail Report",
			is_private=1,
		)

		return {
			"file_url": saved.file_url,
			"filename": filename,
			"total_records": len(data),
		}

	elif download_type == "center_wise":
		# Group by test centre and package into a ZIP
		from collections import OrderedDict
		centre_groups = OrderedDict()
		for d in data:
			c_name = (d.get("test_centre") or "Unassigned Centre").strip()
			centre_groups.setdefault(c_name, []).append(d)

		zip_buf = io.BytesIO()

		with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
			for c_name, c_rows in centre_groups.items():
				wb = openpyxl.Workbook()
				ws = wb.active
				ws.title = "Candidate Attendance"

				_format_candidate_sheet(ws, c_name, c_rows, include_centre=False)

				wb_buf = io.BytesIO()
				wb.save(wb_buf)
				wb_buf.seek(0)

				safe_c = _safe_filename(c_name)
				# Directly write files to zip root (no extra nested subfolder)
				zf.writestr(f"{safe_c}.xlsx", wb_buf.getvalue())

		zip_buf.seek(0)
		zip_filename = f"Test_Center_Wise_Candidate_Details{ay_part}_{ts}.zip"

		saved = save_file(
			zip_filename,
			zip_buf.getvalue(),
			"Report",
			"Test Center-wise Candidate Detail Report",
			is_private=1,
		)

		return {
			"file_url": saved.file_url,
			"filename": zip_filename,
			"total_centres": len(centre_groups),
			"total_records": len(data),
		}

	else:
		frappe.throw(_("Invalid download option specified."))
